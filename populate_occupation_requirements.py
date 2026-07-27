from __future__ import annotations

import copy
import os
from pathlib import Path

import openpyxl
from openpyxl.styles.cell_style import StyleArray


ROOT = Path(r"C:\Users\mudupi\Desktop\Cursor\Moga Dev")
MD_PATH = ROOT / "CV_QDE_Occupation_Details_Requirements.md"
XLSX_PATH = ROOT / "User Stories_CE.xlsx"

STAGE = "QDE"
SUB_STAGE = "Occupation Details"


def field(
    name: str,
    obj: str,
    field_type: str,
    values: str | None,
    section: str,
    validation: str,
    error: str | None,
    control: str | None,
    expression: str | None,
) -> list[str | None]:
    return [
        name,
        obj,
        field_type,
        values,
        STAGE,
        SUB_STAGE,
        section,
        validation,
        error,
        control,
        expression,
    ]


employment_types = "Full-Time; Part-Time; Contract; Intern"
industries = (
    "Agriculture; Automotive; Banking & Finance; Construction; Education; Government; "
    "Healthcare; IT & Technology; Manufacturing; Retail & Trade; Transport & Logistics; Other"
)
professional_bodies = (
    "ICAI; ICSI; Medical Council of India; State Medical Council; Bar Council of India; "
    "State Bar Council; Council of Architecture; Other"
)
file_validation = (
    "Optional under the source workbook and does not currently gate Save or Next. "
    "The current uploader permits files up to 10 MB each, does not enforce a file-count or "
    "extension allow-list, and permits video; final controls require business confirmation."
)

fields: list[list[str | None]] = [
    field(
        "Employer Name",
        "Applicant_Employment__c.Employer_Name__c",
        "Text",
        None,
        "Salaried — Primary Employment",
        "Required for Salaried applicants; manually entered with no employer lookup on the live screen.",
        "Please enter Employer Name.",
        "Applicant Income Category",
        "Visible when Applicant Income Category = Salaried.",
    ),
    field(
        "Job Title / Designation",
        "Applicant_Employment__c.Job_Title_Designation__c",
        "Text",
        None,
        "Salaried — Primary Employment",
        "Optional under the source workbook.",
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Salaried.",
    ),
    field(
        "Employment Type",
        "Applicant_Employment__c.Employment_Type__c",
        "Picklist",
        employment_types,
        "Salaried — Primary Employment",
        "Required for Salaried applicants.",
        "Please select Employment Type.",
        "Applicant Income Category",
        "Visible when Applicant Income Category = Salaried.",
    ),
    field(
        "Industry / Sector",
        "Applicant_Employment__c.Industry_Sector__c",
        "Picklist",
        industries,
        "Salaried — Primary Employment",
        "Optional under the source workbook.",
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Salaried.",
    ),
    field(
        "Employment Start Date",
        "Applicant_Employment__c.Employment_Start_Date__c",
        "Date",
        None,
        "Salaried — Primary Employment",
        "Optional; when entered, cannot be later than today.",
        "Employment Start Date cannot be ahead of today.",
        "Applicant Income Category",
        "Visible when Applicant Income Category = Salaried.",
    ),
    field(
        "Monthly Gross Salary",
        "Applicant_Employment__c.Monthly_Gross_Salary__c",
        "Currency",
        None,
        "Salaried — Primary Employment",
        "Required; the intended value is zero or greater. Server-side enforcement of the lower bound requires alignment.",
        "Please enter Monthly Gross Salary.",
        "Applicant Income Category",
        "Visible when Applicant Income Category = Salaried.",
    ),
    field(
        "Work Address",
        "Applicant_Employment__c.Work_Address__c",
        "Text Area",
        None,
        "Salaried — Primary Employment (source/repository field)",
        "Optional under the source workbook. This combined field exists in the repository but was not displayed on the live screen.",
        None,
        "Applicant Income Category",
        "Source/repository field for Salaried; the live screen instead displays Office Address Details when Office Address Same as Current = No.",
    ),
    field(
        "Work Email Address",
        "Applicant_Employment__c.Work_Email_Address__c",
        "Email with Verify action",
        None,
        "Salaried — Primary Employment (source/repository field)",
        "Optional under the source workbook; valid email format when entered. This primary-employment field exists in the repository but was not displayed on the live screen.",
        "Please enter a valid work email address.",
        "Applicant Income Category",
        "Source/repository field for Salaried; the live screen displays a separate required Office Email when Office Address Same as Current = No.",
    ),
    field(
        "Work Email ID",
        "Applicant_Employment__c.HR_Payroll_Email__c",
        "Email",
        None,
        "Salaried — HR / Payroll Contact",
        "Optional; valid email format when entered.",
        "Please enter a valid HR / payroll work email address.",
        "Applicant Income Category",
        "Visible when Applicant Income Category = Salaried.",
    ),
    field(
        "Direct Phone Number",
        "Applicant_Employment__c.HR_Payroll_Phone_Number__c",
        "Phone",
        None,
        "Salaried — HR / Payroll Contact",
        "Optional; when entered, exactly 10 digits and starts with 6, 7, 8, or 9.",
        "Phone number must be exactly 10 digits; Please enter a valid 10-digit mobile number starting with 6, 7, 8, or 9.",
        "Applicant Income Category",
        "Visible when Applicant Income Category = Salaried.",
    ),
    field(
        "Employer Name",
        "Applicant_Employment__c.Employer_Name__c",
        "Text",
        None,
        "Salaried — Previous Employment",
        "Required for every added previous-employment record, although the live label does not show an asterisk.",
        "Please enter Employer Name.",
        "Applicant Income Category + Add Previous Employer",
        "Visible in an added or edited Previous Employment record for Salaried.",
    ),
    field(
        "Job Title / Designation",
        "Applicant_Employment__c.Job_Title_Designation__c",
        "Text",
        None,
        "Salaried — Previous Employment",
        "Optional under the source workbook.",
        None,
        "Applicant Income Category + Add Previous Employer",
        "Visible in an added or edited Previous Employment record for Salaried.",
    ),
    field(
        "Employment Type",
        "Applicant_Employment__c.Employment_Type__c",
        "Picklist",
        employment_types,
        "Salaried — Previous Employment",
        "Optional under the source workbook.",
        None,
        "Applicant Income Category + Add Previous Employer",
        "Visible in an added or edited Previous Employment record for Salaried.",
    ),
    field(
        "Industry / Sector",
        "Applicant_Employment__c.Industry_Sector__c",
        "Picklist",
        industries,
        "Salaried — Previous Employment",
        "Optional under the source workbook.",
        None,
        "Applicant Income Category + Add Previous Employer",
        "Visible in an added or edited Previous Employment record for Salaried.",
    ),
    field(
        "Employment Start Date",
        "Applicant_Employment__c.Employment_Start_Date__c",
        "Date",
        None,
        "Salaried — Previous Employment",
        "Optional; when entered, cannot be later than today. Start date is not currently checked against End Date.",
        "Previous employment start date cannot be future date.",
        "Applicant Income Category + Add Previous Employer",
        "Visible in an added or edited Previous Employment record for Salaried.",
    ),
    field(
        "Employment End Date",
        "Applicant_Employment__c.Employment_End_Date__c",
        "Date",
        None,
        "Salaried — Previous Employment",
        "Required for every previous-employment record and cannot be later than today.",
        "Please enter Employment End Date for previous employment; Previous employment end date cannot be future date.",
        "Applicant Income Category + Add Previous Employer",
        "Visible in an added or edited Previous Employment record for Salaried.",
    ),
    field(
        "Work Email Address",
        "Applicant_Employment__c.Work_Email_Address__c",
        "Email",
        None,
        "Salaried — Previous Employment",
        "Optional; valid email format when entered.",
        "Please enter a valid work email address for previous employment.",
        "Applicant Income Category + Add Previous Employer",
        "Visible in an added or edited Previous Employment record for Salaried.",
    ),
    field(
        "Phone Number",
        "Applicant_Employment__c.HR_Payroll_Phone_Number__c",
        "Phone",
        None,
        "Salaried — Previous Employment",
        "Optional; when entered, exactly 10 digits and starts with 6, 7, 8, or 9.",
        "Phone number must be exactly 10 digits; Please enter a valid 10-digit mobile number starting with 6, 7, 8, or 9.",
        "Applicant Income Category + Add Previous Employer",
        "Visible in an added or edited Previous Employment record for Salaried.",
    ),
    field(
        "Profession Type",
        "Applicant_Employment__c.Profession_Type__c",
        "Tile selection",
        (
            "Doctor / Medical Professional; Chartered Accountant; Advocate / Lawyer; Architect; "
            "Independent Consultant; Other Self-Employed Professional"
        ),
        "Self-Employed Professional — Primary Employment",
        "Required. Selecting a profession displays the applicable registration and professional-body labels.",
        "Please select a Profession Type.",
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Professional.",
    ),
    field(
        "Practice / Firm Name",
        "Applicant_Employment__c.Practice_Firm_Name__c",
        "Text",
        None,
        "Self-Employed Professional — Primary Employment",
        "Required; manually entered with no practice or firm lookup on the live screen.",
        "Please enter Practice / Firm Name.",
        "Profession Type",
        "Visible after a Profession Type is selected.",
    ),
    field(
        "Medical Council Registration No.",
        "Applicant_Employment__c.Medical_Council_Registration_Number__c",
        "Text",
        None,
        "Self-Employed Professional — Primary Employment",
        "Optional under the source workbook.",
        None,
        "Profession Type",
        "Visible when Profession Type = Doctor / Medical Professional.",
    ),
    field(
        "Medical Council",
        "Applicant_Employment__c.Professional_Body__c",
        "Picklist",
        "Medical Council of India; State Medical Council",
        "Self-Employed Professional — Primary Employment",
        "Optional under the source workbook.",
        None,
        "Profession Type",
        "Visible when Profession Type = Doctor / Medical Professional.",
    ),
    field(
        "ICAI Membership No.",
        "Applicant_Employment__c.ICAI_Membership_No__c",
        "Text",
        None,
        "Self-Employed Professional — Primary Employment (live enrichment)",
        "Optional.",
        None,
        "Profession Type",
        "Visible when Profession Type = Chartered Accountant.",
    ),
    field(
        "Professional Body",
        "Applicant_Employment__c.Professional_Body__c",
        "Picklist",
        professional_bodies,
        "Self-Employed Professional — Primary Employment (live enrichment)",
        "Optional; choices are filtered for regulated professions and use the full list for Independent Consultant and Other.",
        None,
        "Profession Type",
        "Chartered Accountant: ICAI or ICSI; Architect: Council of Architecture; Independent Consultant or Other: full list.",
    ),
    field(
        "Bar Council Enrollment No.",
        "Applicant_Employment__c.Bar_Council_Enrollment_No__c",
        "Text",
        None,
        "Self-Employed Professional — Primary Employment (live enrichment)",
        "Optional.",
        None,
        "Profession Type",
        "Visible when Profession Type = Advocate / Lawyer.",
    ),
    field(
        "Bar Council",
        "Applicant_Employment__c.Bar_Council__c",
        "Picklist",
        "Bar Council of India; State Bar Council",
        "Self-Employed Professional — Primary Employment (live enrichment)",
        "Optional.",
        None,
        "Profession Type",
        "Visible when Profession Type = Advocate / Lawyer.",
    ),
    field(
        "Council of Architecture Reg. No.",
        "Applicant_Employment__c.Council_of_Architecture_Reg_No__c",
        "Text",
        None,
        "Self-Employed Professional — Primary Employment (live enrichment)",
        "Optional.",
        None,
        "Profession Type",
        "Visible when Profession Type = Architect.",
    ),
    field(
        "Professional Certification / Licence No.",
        "Applicant_Employment__c.Professional_Certification_Licence_No__c",
        "Text",
        None,
        "Self-Employed Professional — Primary Employment (live enrichment)",
        "Optional.",
        None,
        "Profession Type",
        "Visible when Profession Type = Independent Consultant.",
    ),
    field(
        "Registration / Licence No.",
        "Applicant_Employment__c.Registration_Licence_No__c",
        "Text",
        None,
        "Self-Employed Professional — Primary Employment (live enrichment)",
        "Optional.",
        None,
        "Profession Type",
        "Visible when Profession Type = Other Self-Employed Professional.",
    ),
    field(
        "Date of Registration",
        "Applicant_Employment__c.Date_of_Registration__c",
        "Date",
        None,
        "Self-Employed Professional — Primary Employment",
        "Optional under the source workbook; no past-date rule is currently enforced.",
        None,
        "Profession Type",
        "Visible after a Profession Type is selected.",
    ),
    field(
        "Years in Practice",
        "Applicant_Employment__c.Years_in_Practice__c",
        "Number",
        None,
        "Self-Employed Professional — Primary Employment",
        "Optional; intended minimum is 0, but server-side enforcement requires alignment.",
        None,
        "Profession Type",
        "Visible after a Profession Type is selected.",
    ),
    field(
        "No. of Employees / Staff",
        "Applicant_Employment__c.Number_of_Employees_Staff__c",
        "Number",
        None,
        "Self-Employed Professional — Primary Employment",
        "Optional; intended minimum is 0, but server-side enforcement requires alignment.",
        None,
        "Profession Type",
        "Visible after a Profession Type is selected.",
    ),
    field(
        "Annual Professional Income",
        "Applicant_Employment__c.Annual_Professional_Income__c",
        "Currency",
        None,
        "Self-Employed Professional — Primary Employment",
        "Required; intended value is zero or greater. Server-side enforcement of the lower bound requires alignment.",
        "Please enter Annual Professional Income.",
        "Profession Type",
        "Visible after a Profession Type is selected.",
    ),
    field(
        "PAN Number",
        "Applicant_Employment__c.PAN_Number__c",
        "Text",
        None,
        "Self-Employed Professional — Primary Employment",
        "Optional under the source workbook; no PAN-format validation is currently enforced.",
        None,
        "Profession Type",
        "Visible after a Profession Type is selected.",
    ),
    field(
        "GST Registration No. (if applicable)",
        "Applicant_Employment__c.GST_Registration_No__c",
        "Text",
        None,
        "Self-Employed Professional — Primary Employment",
        "Optional under the source workbook; no GST-format validation is currently enforced.",
        None,
        "Profession Type",
        "Visible after a Profession Type is selected.",
    ),
    field(
        "Practice / Office Address",
        "Applicant_Employment__c.Work_Address__c",
        "Text Area",
        None,
        "Self-Employed Professional — Primary Employment (source/repository field)",
        "Optional under the source workbook. This combined field exists in the repository but was not displayed on the live screen.",
        None,
        "Applicant Income Category",
        "Source/repository field; the live screen instead displays Office Address Details when Office Address Same as Current = No.",
    ),
    field(
        "Upload Registration Certificate",
        "Document_CheckList__c.ContentVersionId__c",
        "File upload",
        "Current implementation has no extension allow-list",
        "Self-Employed Professional — Supporting Documents",
        file_validation,
        None,
        "Applicant Income Category + Profession Type",
        "Visible after a Self-Employed Professional Profession Type is selected.",
    ),
    field(
        "Upload Practice Licence",
        "Document_CheckList__c.ContentVersionId__c",
        "File upload",
        "Current implementation has no extension allow-list",
        "Self-Employed Professional — Supporting Documents",
        file_validation,
        None,
        "Applicant Income Category + Profession Type",
        "Visible after a Self-Employed Professional Profession Type is selected.",
    ),
    field(
        "Upload Latest ITR (last 3 years)",
        "Document_CheckList__c.ContentVersionId__c",
        "File upload",
        "Current implementation has no extension allow-list",
        "Self-Employed Professional — Supporting Documents",
        file_validation,
        None,
        "Applicant Income Category + Profession Type",
        "Visible after a Self-Employed Professional Profession Type is selected.",
    ),
    field(
        "Business / Firm Name",
        "Applicant_Employment__c.Business_Firm_Name__c",
        "Text",
        None,
        "Self-Employed Non-Professional — Primary Employment",
        "Required; manually entered with no business lookup on the live screen.",
        "Please enter Business / Firm Name.",
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "Nature of Business",
        "Applicant_Employment__c.Nature_of_Business__c",
        "Picklist",
        "Trading; Manufacturing; Retail; Wholesale; Services; Agriculture; Other",
        "Self-Employed Non-Professional — Primary Employment",
        "Optional under the source workbook.",
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "Business Constitution",
        "Applicant_Employment__c.Business_Constitution__c",
        "Picklist",
        "Proprietorship; Partnership; Private Limited; Public Limited; LLP; Other",
        "Self-Employed Non-Professional — Primary Employment",
        "Optional under the source workbook.",
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "Date of Commencement",
        "Applicant_Employment__c.Date_of_Commencement__c",
        "Date",
        None,
        "Self-Employed Non-Professional — Primary Employment",
        "Optional; when entered, cannot be later than today.",
        "Date of Commencement cannot be ahead of today.",
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "Business Vintage (Years)",
        "Applicant_Employment__c.Business_Vintage_Years__c",
        "Number",
        None,
        "Self-Employed Non-Professional — Primary Employment",
        "Optional; intended minimum is 0, but server-side enforcement requires alignment.",
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "No. of Employees / Staff",
        "Applicant_Employment__c.Number_of_Employees_Staff__c",
        "Number",
        None,
        "Self-Employed Non-Professional — Primary Employment",
        "Optional; intended minimum is 0, but server-side enforcement requires alignment.",
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "Annual Turnover",
        "Applicant_Employment__c.Annual_Turnover__c",
        "Currency",
        None,
        "Self-Employed Non-Professional — Primary Employment",
        "Required; intended value is zero or greater. Server-side enforcement of the lower bound requires alignment.",
        "Please enter Annual Turnover.",
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "Annual Net Income / Profit",
        "Applicant_Employment__c.Annual_Net_Income__c",
        "Currency",
        None,
        "Self-Employed Non-Professional — Primary Employment",
        "Optional; intended minimum is 0, but server-side enforcement requires alignment.",
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "Business / Proprietor PAN",
        "Applicant_Employment__c.PAN_Number__c",
        "Text",
        None,
        "Self-Employed Non-Professional — Primary Employment",
        "Optional under the source workbook; no PAN-format validation is currently enforced.",
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "GST Registration No. (if applicable)",
        "Applicant_Employment__c.GST_Registration_No__c",
        "Text",
        None,
        "Self-Employed Non-Professional — Primary Employment",
        "Optional under the source workbook; no GST-format validation is currently enforced.",
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "Udyam / MSME Reg. No. (if applicable)",
        "Applicant_Employment__c.Udyam_MSME_Reg_No__c",
        "Text",
        None,
        "Self-Employed Non-Professional — Primary Employment",
        "Optional under the source workbook; no Udyam-format validation is currently enforced.",
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "Business / Shop Address",
        "Applicant_Employment__c.Business_Shop_Address__c",
        "Text Area",
        None,
        "Self-Employed Non-Professional — Primary Employment (source/repository field)",
        "Optional under the source workbook. This combined field exists in the repository but was not displayed on the live screen.",
        None,
        "Applicant Income Category",
        "Source/repository field; the live screen instead displays Office Address Details when Office Address Same as Current = No.",
    ),
    field(
        "Upload Business Registration / Shop Act Licence",
        "Document_CheckList__c.ContentVersionId__c",
        "File upload",
        "Current implementation has no extension allow-list",
        "Self-Employed Non-Professional — Supporting Documents",
        file_validation,
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "Upload GST Returns (last 12 months)",
        "Document_CheckList__c.ContentVersionId__c",
        "File upload",
        "Current implementation has no extension allow-list",
        "Self-Employed Non-Professional — Supporting Documents",
        file_validation,
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "Upload Latest ITR (last 3 years)",
        "Document_CheckList__c.ContentVersionId__c",
        "File upload",
        "Current implementation has no extension allow-list",
        "Self-Employed Non-Professional — Supporting Documents",
        file_validation,
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "Upload Bank Statements (last 6 months)",
        "Document_CheckList__c.ContentVersionId__c",
        "File upload",
        "Current implementation has no extension allow-list",
        "Self-Employed Non-Professional — Supporting Documents",
        file_validation,
        None,
        "Applicant Income Category",
        "Visible when Applicant Income Category = Self Employed Non Professional.",
    ),
    field(
        "Is Office Address Same As Current Address?",
        "Applicant_Address__c.Is_Office_Address_Same_As_Current__c",
        "Segmented Yes / No",
        "Yes; No",
        "All earning profiles — Primary Employment",
        "Required on the live screen. Yes hides Office Address Details; No displays them.",
        "Select whether the office address is the same as the current address.",
        "Applicant Income Category",
        "Visible for Salaried, Self Employed Professional, and Self Employed Non Professional.",
    ),
    field(
        "Office Address Line 1",
        "Applicant_Address__c.Address_Line_1__c",
        "Text",
        None,
        "All earning profiles — Office Address Details",
        "Required; maximum 255 characters.",
        "Office Address is incomplete. Missing: Address Line 1.",
        "Is Office Address Same As Current Address?",
        "Visible and required when answer = No.",
    ),
    field(
        "Office Address Line 2",
        "Applicant_Address__c.Address_Line_2__c",
        "Text",
        None,
        "All earning profiles — Office Address Details",
        "Optional; maximum 255 characters.",
        None,
        "Is Office Address Same As Current Address?",
        "Visible when answer = No.",
    ),
    field(
        "Office Address Line 3",
        "Applicant_Address__c.Address_Line_3__c",
        "Text",
        None,
        "All earning profiles — Office Address Details",
        "Optional; maximum 255 characters.",
        None,
        "Is Office Address Same As Current Address?",
        "Visible when answer = No.",
    ),
    field(
        "Office PIN Code",
        "Applicant_Address__c.PinCode__c",
        "Numeric text",
        None,
        "All earning profiles — Office Address Details",
        "Required; exactly 6 digits. Entering a valid PIN triggers City and State lookup.",
        "Office Address is incomplete. Missing: Pin Code; Please wait — looking up city and state for the pin code.",
        "Is Office Address Same As Current Address?",
        "Visible and required when answer = No.",
    ),
    field(
        "Office City",
        "Applicant_Address__c.City__c",
        "Text",
        None,
        "All earning profiles — Office Address Details",
        "Required; read-only when the PIN lookup returns a value and editable when no value is returned.",
        "Office Address is incomplete. Missing: City.",
        "Office PIN Code",
        "Visible when Office Address Same as Current = No; populated from PIN lookup when available.",
    ),
    field(
        "Office State",
        "Applicant_Address__c.State__c",
        "Text",
        None,
        "All earning profiles — Office Address Details",
        "Required; read-only when the PIN lookup returns a value and editable when no value is returned.",
        "Office Address is incomplete. Missing: State.",
        "Office PIN Code",
        "Visible when Office Address Same as Current = No; populated from PIN lookup when available.",
    ),
    field(
        "Office Work Email Address",
        "Applicant_Address__c.Email__c",
        "Email with Send OTP action",
        None,
        "All earning profiles — Office Address Details",
        "Required when Office Address Same as Current = No; valid email format and successful OTP verification are required.",
        "Please enter a valid Office Email address; Please verify the Office Email before continuing.",
        "Is Office Address Same As Current Address?",
        "Visible when answer = No. Changing a verified email requires re-verification.",
    ),
]


stories = [
    {
        "id": "US-QDE-OCCUP-CV-001",
        "title": "Route the Applicant to the Applicable Occupation Profile",
        "description": (
            "As a Sales Officer, I want the Occupation Details screen to display the profile that matches "
            "the applicant's income category so that only relevant employment or business information is captured."
        ),
        "scenarios": [
            (
                "Positive Scenario — Earning profiles",
                [
                    "Given the applicant is Individual and has an earning income category",
                    "When Occupation Details is opened",
                    "Then exactly one of Salaried, Self-Employed Professional, or Self-Employed Non-Professional is active",
                    "And the other profile tabs are disabled",
                ],
            ),
            (
                "Positive Scenario — Non-Earning",
                [
                    "Given the Individual applicant is Non-Earning",
                    "When Occupation Details is reached",
                    "Then a not-applicable state is displayed and employment capture is not required",
                ],
            ),
            (
                "Navigation Scenario — Non-Individual",
                [
                    "Given the applicant is Non-Individual",
                    "When the QDE journey is evaluated",
                    "Then Occupation Details is skipped and the applicant is routed to Dedupe and Eligibility",
                ],
            ),
            (
                "Validation Scenario",
                [
                    "Given the income category is blank or unsupported",
                    "When Occupation Details is opened",
                    "Then the system must not silently default the applicant to Salaried",
                    "And the Sales Officer is instructed to complete or correct the income category",
                ],
            ),
            (
                "Edge Case Scenario",
                [
                    "Given a Non-Individual application contains an Individual guarantor or beneficiary",
                    "When that Individual applicant's QDE journey is opened",
                    "Then Occupation Details is available according to that applicant's own income category",
                ],
            ),
        ],
        "assumptions": [
            "Income category is captured and approved in Basic Details before Occupation Details.",
            "The current mandatory status in the supplied occupation workbook is authoritative until Mandatory Final is populated.",
        ],
        "integrations": [
            "No external integration is triggered by profile routing.",
            "The QDE stage-routing service supplies the applicable applicant type, income category, and next step.",
        ],
        "scope": [
            "Profile selection, disabled tabs, Non-Earning handling, and Non-Individual skip behavior.",
        ],
    },
    {
        "id": "US-QDE-OCCUP-CV-002",
        "title": "Capture Salaried Primary Employment Details",
        "description": (
            "As a Sales Officer, I want to capture the applicant's current salaried employment and income "
            "so that the application contains the information required for employment assessment."
        ),
        "scenarios": [
            (
                "Positive Scenario",
                [
                    "Given the applicant income category is Salaried",
                    "When Primary Employment is displayed",
                    "Then Employer Name, Employment Type, and Monthly Gross Salary are required",
                    "And Job Title, Industry, Employment Start Date, Work Address, and Work Email are optional under the supplied workbook",
                ],
            ),
            (
                "Validation Scenario",
                [
                    "Given a required Salaried value is missing",
                    "When Save is selected",
                    "Then saving is blocked and the first missing required value is identified",
                    "And Employment Start Date, when entered, cannot be later than today",
                ],
            ),
            (
                "Lookup Scenario",
                [
                    "Given the Sales Officer enters Employer Name",
                    "When the value is captured",
                    "Then it is accepted as manual text because no external employer lookup is currently configured",
                    "And Industry and Employment Type are selected from their controlled lists",
                ],
            ),
            (
                "Persistence Scenario",
                [
                    "Given valid Salaried details are saved",
                    "When the screen is revisited",
                    "Then the saved current-employment values are displayed",
                ],
            ),
        ],
        "assumptions": [
            "Monthly Gross Salary represents the applicant's current gross monthly salary.",
            "Optional source-workbook fields remain available even when they do not gate Save.",
        ],
        "integrations": [
            "No external employer verification or employer-name lookup is currently invoked.",
            "Controlled Employment Type and Industry values are supplied by the application.",
        ],
        "scope": [
            "Current Salaried employment, income, date, industry, and optional work-contact details.",
        ],
    },
    {
        "id": "US-QDE-OCCUP-CV-003",
        "title": "Capture and Verify a Different Office Address",
        "description": (
            "As a Sales Officer, I want to confirm whether the office address matches the current address "
            "and capture a different office address when required so that work-location and contact information is verified."
        ),
        "scenarios": [
            (
                "Positive Scenario — Same address",
                [
                    "Given the office address is the same as the current address",
                    "When Yes is selected",
                    "Then Office Address Details is hidden and separate address entry is not required",
                ],
            ),
            (
                "Positive Scenario — Different address",
                [
                    "Given the office address is different",
                    "When No is selected",
                    "Then Address Line 1, PIN Code, City, State, and Office Work Email are required",
                    "And Address Lines 2 and 3 remain optional",
                ],
            ),
            (
                "Integration Scenario — PIN lookup",
                [
                    "Given a valid six-digit PIN is entered",
                    "When the PIN lookup succeeds",
                    "Then City and State are populated and protected from accidental editing",
                    "And if the lookup fails or returns no value, City and State remain available for manual entry",
                ],
            ),
            (
                "Integration Scenario — Office email OTP",
                [
                    "Given a valid Office Work Email is entered",
                    "When Send OTP is selected",
                    "Then an OTP is sent and successful verification marks the email as verified",
                    "And a failed, expired, or incorrect OTP blocks completion while allowing an approved resend or retry",
                ],
            ),
            (
                "Validation Scenario",
                [
                    "Given Office Address Same as Current is No",
                    "When any required address value is missing, the PIN lookup is still running, the email format is invalid, or email verification is incomplete",
                    "Then Save or Next is blocked with a clear reason",
                ],
            ),
            (
                "Edge Case Scenario — Toggle retention",
                [
                    "Given office-address values have been entered under No",
                    "When the Sales Officer changes the answer to Yes and then back to No before saving",
                    "Then the current live behavior retains the entered values",
                    "And the client must confirm whether retained or cleared behavior is required",
                ],
            ),
        ],
        "assumptions": [
            "The office-address question applies to all three earning occupation profiles.",
            "A verified Office Work Email is mandatory only when a separate office address is captured.",
        ],
        "integrations": [
            "PIN Code lookup supplies City, State, and a system-derived Country where available.",
            "Office Email OTP is triggered by Send OTP; failure, expiry, resend, and retry must not mark the email verified.",
        ],
        "scope": [
            "Office-address dependency, PIN lookup, Office Work Email validation, and OTP verification.",
        ],
    },
    {
        "id": "US-QDE-OCCUP-CV-004",
        "title": "Capture Salaried HR Contact and Previous Employment",
        "description": (
            "As a Sales Officer, I want to capture optional HR contact information and the applicant's prior "
            "employment so that employment history is available for assessment."
        ),
        "scenarios": [
            (
                "Positive Scenario — HR contact",
                [
                    "Given the Salaried profile is active",
                    "When HR or Payroll contact information is available",
                    "Then Work Email ID and Direct Phone Number can be entered",
                ],
            ),
            (
                "Positive Scenario — Previous employment",
                [
                    "Given the Sales Officer selects Add Previous Employer",
                    "When a history record is entered",
                    "Then Employer Name and Employment End Date are required",
                    "And Job Title, Employment Type, Industry, Start Date, Work Email, and Phone remain optional",
                ],
            ),
            (
                "Validation Scenario",
                [
                    "Given an optional email or phone is entered",
                    "When the value is validated",
                    "Then the email must use a valid format and the phone must contain exactly 10 digits beginning with 6, 7, 8, or 9",
                    "And previous-employment dates cannot be later than today",
                ],
            ),
            (
                "Edit and Delete Scenario",
                [
                    "Given a previous-employment record exists",
                    "When it is edited, deleted, or cancelled",
                    "Then the list reflects the intended action and unrelated records remain unchanged",
                ],
            ),
            (
                "Edge Case Scenario",
                [
                    "Given Employment Start Date is after Employment End Date or the record falls outside the intended history period",
                    "When the record is validated",
                    "Then the signed-off chronology and history-window rule must be applied once confirmed",
                ],
            ),
        ],
        "assumptions": [
            "Previous Employment applies only to the Salaried profile.",
            "The business will confirm whether 'last 3 years' is a date boundary, a record-count rule, or descriptive guidance.",
        ],
        "integrations": [
            "No external HR, payroll, employer, email, or phone verification integration is currently invoked by these fields.",
        ],
        "scope": [
            "HR contact validation and add, edit, delete, cancel, and persistence behavior for previous employment.",
        ],
    },
    {
        "id": "US-QDE-OCCUP-CV-005",
        "title": "Capture Self-Employed Professional Details",
        "description": (
            "As a Sales Officer, I want to capture profession-specific practice, registration, and income details "
            "so that a self-employed professional's occupation can be assessed."
        ),
        "scenarios": [
            (
                "Positive Scenario — Profession selection",
                [
                    "Given the applicant is Self-Employed Professional",
                    "When a Profession Type is selected",
                    "Then the applicable registration label and professional-body choices are displayed",
                    "And Practice or Firm Name and Annual Professional Income are required",
                ],
            ),
            (
                "Dependency Scenario",
                [
                    "Given Doctor, Chartered Accountant, Advocate, Architect, Independent Consultant, or Other is selected",
                    "When the profile form is displayed",
                    "Then the corresponding registration or licence field is shown",
                    "And only the professional-body choices applicable to that profession are offered where a controlled list exists",
                ],
            ),
            (
                "Validation Scenario",
                [
                    "Given Profession Type, Practice or Firm Name, or Annual Professional Income is missing",
                    "When Save is selected",
                    "Then saving is blocked and the missing value is identified",
                    "And the signed-off non-negative and identifier-format rules are applied when confirmed",
                ],
            ),
            (
                "Lookup Scenario",
                [
                    "Given a practice, firm, registration number, or professional body is entered",
                    "When the value is captured",
                    "Then names and registration numbers are manually entered and professional bodies use controlled lists",
                    "And no external professional-register verification is currently invoked",
                ],
            ),
            (
                "Edge Case Scenario — Profession change",
                [
                    "Given values were entered for one profession",
                    "When Profession Type changes",
                    "Then fields that no longer apply are hidden",
                    "And the client must confirm whether their hidden values are cleared or retained",
                ],
            ),
            (
                "Persistence Scenario",
                [
                    "Given professional details are saved",
                    "When the screen is revisited",
                    "Then the selected profession and saved applicable values are displayed",
                ],
            ),
        ],
        "assumptions": [
            "Profession-specific registration and body fields remain optional because the source workbook marks the supplied registration fields as optional.",
            "Annual Professional Income is the annual amount used for occupation assessment.",
        ],
        "integrations": [
            "No external medical, accounting, legal, architecture, or professional-register lookup is currently invoked.",
            "Professional-body options are controlled by the selected Profession Type.",
        ],
        "scope": [
            "All six profession types, their labels, controlled values, required fields, and profile persistence.",
        ],
    },
    {
        "id": "US-QDE-OCCUP-CV-006",
        "title": "Capture Self-Employed Non-Professional Business Details",
        "description": (
            "As a Sales Officer, I want to capture business identity, scale, and income details so that a "
            "self-employed non-professional applicant's occupation can be assessed."
        ),
        "scenarios": [
            (
                "Positive Scenario",
                [
                    "Given the applicant is Self-Employed Non-Professional",
                    "When Primary Employment is displayed",
                    "Then Business or Firm Name and Annual Turnover are required",
                    "And the remaining source-workbook business fields are optional",
                ],
            ),
            (
                "Controlled Value Scenario",
                [
                    "Given Nature of Business or Business Constitution is entered",
                    "When a choice is selected",
                    "Then the value comes from its complete controlled list",
                ],
            ),
            (
                "Validation Scenario",
                [
                    "Given a required business value is missing",
                    "When Save is selected",
                    "Then saving is blocked and the missing value is identified",
                    "And Date of Commencement cannot be later than today",
                ],
            ),
            (
                "Lookup Scenario",
                [
                    "Given the Sales Officer enters a business name or business identifier",
                    "When the value is captured",
                    "Then it is accepted as manual input because no external business, PAN, GST, or Udyam lookup is currently configured",
                ],
            ),
            (
                "Persistence Scenario",
                [
                    "Given business details are saved",
                    "When the screen is revisited",
                    "Then the saved business, income, and identifier values are displayed",
                ],
            ),
        ],
        "assumptions": [
            "Annual Turnover is mandatory and Annual Net Income remains optional under the source workbook.",
            "Business identifier formats will be enforced only after the client confirms the required rules.",
        ],
        "integrations": [
            "No external business-name, PAN, GST, Udyam, or company-register lookup is currently invoked.",
            "Nature of Business and Business Constitution values are supplied by controlled application lists.",
        ],
        "scope": [
            "Business profile fields, controlled values, validation, and persistence.",
        ],
    },
    {
        "id": "US-QDE-OCCUP-CV-007",
        "title": "Upload Occupation Supporting Documents",
        "description": (
            "As a Sales Officer, I want to upload occupation documents for self-employed applicants so that "
            "supporting evidence is available for verification."
        ),
        "scenarios": [
            (
                "Positive Scenario — Professional",
                [
                    "Given a Profession Type is selected",
                    "When Supporting Documents is displayed",
                    "Then Registration Certificate, Practice Licence, and Latest ITR slots are available",
                ],
            ),
            (
                "Positive Scenario — Non-Professional",
                [
                    "Given the applicant is Self-Employed Non-Professional",
                    "When Supporting Documents is displayed",
                    "Then Business Registration, GST Returns, Latest ITR, and Bank Statement slots are available",
                ],
            ),
            (
                "Upload Scenario",
                [
                    "Given a file is selected",
                    "When upload succeeds",
                    "Then the file can be previewed and removed and remains available after the screen is revisited",
                ],
            ),
            (
                "Failure and Retry Scenario",
                [
                    "Given an upload, preview, or delete action fails",
                    "When the failure is returned",
                    "Then the Sales Officer sees a clear message, entered occupation data is retained, and the action can be retried",
                ],
            ),
            (
                "Validation Scenario",
                [
                    "Given document rules are evaluated",
                    "When Save or Next is selected",
                    "Then the current optional document slots do not block progression",
                    "And approved file types, file counts, size limits, and any mandatory rules must be enforced once signed off",
                ],
            ),
        ],
        "assumptions": [
            "All occupation document slots remain optional until the client approves final mandatory rules.",
            "Only active, successfully uploaded documents are shown on revisit.",
        ],
        "integrations": [
            "The document upload service stores and retrieves occupation documents.",
            "Upload success, failure, preview, removal, and retry behavior are part of this integration.",
        ],
        "scope": [
            "Professional and business document slots, upload, preview, remove, retry, and persistence.",
        ],
    },
    {
        "id": "US-QDE-OCCUP-CV-008",
        "title": "Save Occupation Details and Continue to Dedupe",
        "description": (
            "As a Sales Officer, I want to save valid Occupation Details and continue to Dedupe and Eligibility "
            "so that the applicant progresses with complete occupation information."
        ),
        "scenarios": [
            (
                "Positive Scenario — Save",
                [
                    "Given all mandatory fields for the active occupation profile are complete",
                    "When Save is selected",
                    "Then current and previous employment details and applicable office information are retained",
                    "And Next becomes available only after a successful save",
                ],
            ),
            (
                "Validation Scenario",
                [
                    "Given required data is missing or invalid",
                    "When Save is selected",
                    "Then saving is blocked and the Sales Officer is told what must be corrected",
                    "And Next remains unavailable",
                ],
            ),
            (
                "Change Scenario",
                [
                    "Given details were saved and then changed",
                    "When the screen detects unsaved changes",
                    "Then Next is disabled until the revised values are saved successfully",
                ],
            ),
            (
                "Navigation Scenario",
                [
                    "Given valid details have been saved",
                    "When Next is selected",
                    "Then Occupation Details is marked complete and Dedupe and Eligibility is opened",
                ],
            ),
            (
                "Failure Scenario",
                [
                    "Given save, completion, or navigation fails",
                    "When the failure is returned",
                    "Then the Sales Officer receives a clear message, entered data remains available, and duplicate completion is avoided",
                ],
            ),
            (
                "Review Scenario",
                [
                    "Given Occupation Details is complete",
                    "When the Sales Officer revisits the step",
                    "Then saved values and documents are displayed in read-only form unless an approved reopen process is used",
                ],
            ),
            (
                "Edge Case Scenario — Profile change",
                [
                    "Given saved occupation data exists and the income category changes",
                    "When the new profile is opened",
                    "Then non-applicable values do not influence completion or downstream assessment",
                    "And the client-approved clear or retain rule is applied consistently",
                ],
            ),
        ],
        "assumptions": [
            "Occupation Details completion is required before Dedupe for every earning Individual applicant.",
            "A completed step is read-only unless the business provides an explicit reopen process.",
        ],
        "integrations": [
            "The occupation save service retains current employment, previous employment, and applicable office information.",
            "The QDE stage-routing service opens Dedupe and Eligibility after successful completion.",
        ],
        "scope": [
            "Save gating, unsaved changes, Next, navigation, error recovery, persistence, and completed-stage review.",
        ],
    },
]


discrepancies = [
    "Non-Earning routing conflict: the reference journey requires a not-applicable state, while the reviewed occupation component can treat an unsupported category as Salaried.",
    "Live/repository office-address conflict: the live screen requires an Office Address Same as Current response and displays granular Office Address Details for No; the reviewed parent occupation component instead contains combined work or business address fields.",
    "Source-workbook address conflict: the workbook supplies Work Address, Practice / Office Address, and Business / Shop Address as combined optional fields, while the live screen uses Address Lines, PIN, City, State, and a verified Office Work Email.",
    "Office toggle retention: the live screen retains entered Office Address Details when the answer is changed from No to Yes and back to No before saving. The required clear-or-retain rule is not signed off.",
    "Client/server validation gap: category mandatory fields and non-negative numeric limits are mainly checked in the browser and are not independently enforced before completion.",
    "Completion sequencing gap: occupation completion and navigation are separate operations, so navigation failure can leave the step marked complete.",
    "Immediate read-only behavior: a successful Save makes the current screen read-only before Next, which may prevent correction without a reload or reopen process.",
    "Document control gap: configured file types and maximum file counts are not applied by the current uploader; video is enabled and documents do not gate Save or Next.",
    "Multiple-file persistence gap: multiple files may appear during the session, while only one file reference per slot is reliably restored.",
    "Previous-employment rule gap: start-before-end chronology and the stated three-year history period are not enforced.",
    "Identifier validation gap: PAN, GST, and Udyam values have no implemented format validation in Occupation Details.",
    "Profile/profession reset gap: changing an income category or profession can retain hidden values that no longer apply.",
    "Work-email verification gap: the reviewed OTP flow can alter the applicant's personal email status and does not reliably compare the entered OTP with the issued OTP.",
    "Server completion gap: the completion operation can be invoked without independently proving that a valid occupation record exists.",
    "Additional-applicant gating gap: downstream Dedupe can treat KYC completion as sufficient for an additional Individual applicant even when an earning profile requires Occupation Details.",
    "Test-coverage gap: no direct automated coverage was found for the occupation controller, profile branches, OTP, documents, Save/Next, or failure recovery.",
]

questions = [
    "Should Non-Earning skip Occupation Details entirely or display a not-applicable page with Next?",
    "Please confirm that Currently Mandatory? remains authoritative while Mandatory Final is blank.",
    "Should Office Address Details apply to all three earning profiles exactly as observed on the live screen?",
    "When Office Address Same as Current changes to Yes, should previously entered separate-office values be cleared or retained?",
    "Should the source workbook's combined address fields be replaced by the live granular Office Address Details fields?",
    "Must Office Work Email verification gate Save, Next, or both when a separate office address is entered?",
    "What OTP expiry, resend delay, resend count, attempt count, and retry behavior must apply to Office Work Email?",
    "What exact file types, maximum file count, maximum size, and mandatory rules apply to each occupation document slot, and must video be prohibited?",
    "Should Employer Name, Practice / Firm Name, Business / Firm Name, PAN, GST, Udyam, and professional registration values use external lookup or verification services?",
    "What non-negative, maximum-value, decimal-scale, PAN, GST, Udyam, and registration-number validation rules are required?",
    "Does 'Previous Employment for the last 3 years' mean a date boundary, a maximum number of records, or descriptive guidance?",
    "Should changing income category or Profession Type clear values that no longer apply?",
    "Should completed Occupation Details remain read-only, or should an authorized reopen/edit process be provided?",
    "Should successful completion require the server to revalidate all signed-off mandatory fields and documents before navigation?",
    "Must Dedupe require completed Occupation Details for every additional Individual applicant with an earning income category?",
]


def md_cell(value: str | None) -> str:
    if value is None:
        return ""
    return str(value).replace("|", r"\|").replace("\n", "<br>")


def bullets(items: list[str]) -> str:
    return "\n".join(f"- {item}" for item in items)


def acceptance_text(story: dict) -> str:
    lines = ["Scenarios"]
    for label, criteria in story["scenarios"]:
        lines.append(label)
        lines.extend(f"  • {item}" for item in criteria)
    lines.extend(["", "Assumptions"])
    lines.extend(f"• {item}" for item in story["assumptions"])
    lines.extend(["", "Integrations"])
    lines.extend(f"• {item}" for item in story["integrations"])
    lines.extend(["", "In Scope"])
    lines.extend(f"• {item}" for item in story["scope"])
    return "\n".join(lines)


def story_markdown(story: dict) -> str:
    lines = [
        f"### {story['id']}: {story['title']}",
        "",
        f"**User Story Description:** {story['description']}",
        "",
        "**Acceptance Criteria:**",
        "",
    ]
    for label, criteria in story["scenarios"]:
        lines.append(f"- **{label}:**")
        lines.extend(f"  - {item}" for item in criteria)
    lines.extend(
        [
            "",
            "**Assumptions:**",
            "",
            bullets(story["assumptions"]),
            "",
            "**Integrations:**",
            "",
            bullets(story["integrations"]),
            "",
            "**In Scope:**",
            "",
            bullets(story["scope"]),
            "",
        ]
    )
    return "\n".join(lines)


def build_markdown() -> str:
    headers = [
        "Field Name",
        "Object",
        "Field Type",
        "Drop down values",
        "Display Stage",
        "Display Sub Stage",
        "Display Section",
        "Validation",
        "Error Message",
        "Controlling field",
        "Controlling Expression",
    ]
    table = [
        "| " + " | ".join(headers) + " |",
        "|" + "|".join(["---"] * len(headers)) + "|",
    ]
    table.extend("| " + " | ".join(md_cell(v) for v in row) + " |" for row in fields)

    matrix = [
        "| Rule | Functional behavior |",
        "|---|---|",
        "| Profile routing | The applicant's income category activates one earning profile; other profile tabs are disabled. |",
        "| Salaried mandatory fields | Employer Name, Employment Type, and Monthly Gross Salary. |",
        "| Previous employment mandatory fields | Employer Name and Employment End Date for every added record. |",
        "| Professional mandatory fields | Profession Type, Practice / Firm Name, and Annual Professional Income. |",
        "| Non-Professional mandatory fields | Business / Firm Name and Annual Turnover. |",
        "| Different office address | Address Line 1, six-digit PIN, City, State, and verified Office Work Email are required; Lines 2 and 3 are optional. |",
        "| Date rules | Employment Start Date, previous-employment dates, and Date of Commencement cannot be later than today. |",
        "| Contact rules | Optional emails require valid format; optional phones require 10 digits beginning with 6–9. |",
        "| Save and Next | Next remains disabled until a successful Save; changes after Save require another successful Save. |",
        "| Documents | Current source-workbook document slots are optional and do not gate Save or Next. |",
    ]

    sections = [
        "# CV QDE Occupation Details Requirements",
        "",
        "## Document Control",
        "",
        "- **Stage:** QDE",
        "- **Sub Stage:** Occupation Details",
        "- **Page/Screen Name:** CV QDE Occupation Details",
        "- **Page Key:** QDE-OCCUP-CV",
        "- **Persona:** Sales Officer",
        "- **Applicability used for Excel:** CE = Yes, CV = Yes, Auto = Yes",
        "- **Workbook update mode:** REPLACE_TARGET_PAGE_ROWS",
        "",
        "## Scope and Evidence",
        "",
        "This document combines the supplied mandatory-field workbook, the live Salesforce screen, the reference requirements, and the reviewed occupation implementation. It covers Salaried, Self-Employed Professional, and Self-Employed Non-Professional profiles, including conditional Office Address Details, previous employment, documents, validation, Save, Next, navigation, persistence, and completed-stage review.",
        "",
        "The supplied workbook's **Currently Mandatory?** column is used as the mandatory baseline because **Mandatory Final** is blank on all three source worksheets. Supplied rows are retained. Live-only profession-specific and Office Address Details fields are added and clearly identified through their sections and controlling expressions.",
        "",
        "## Observed Screen Structure",
        "",
        "- The applicant's income category determines the active occupation profile; the other two profile tabs are disabled.",
        "- Salaried displays Primary Employment, HR / Payroll Contact, and Previous Employment.",
        "- Self-Employed Professional first displays six Profession Type tiles, then profession-specific fields and three supporting-document slots.",
        "- Self-Employed Non-Professional displays business fields and four supporting-document slots.",
        "- All three earning profiles display a required Office Address Same as Current choice. Selecting No displays granular Office Address Details and Office Work Email verification.",
        "- Save validates the active profile. Next remains disabled until a successful Save and routes the applicant to Dedupe and Eligibility.",
        "",
        "## Excel Field Details Table",
        "",
        *table,
        "",
        "## Validation and Dependency Matrix",
        "",
        *matrix,
        "",
        "## User Stories",
        "",
        "\n".join(story_markdown(story) for story in stories),
        "## Implementation Discrepancies Requiring Resolution",
        "",
        "\n".join(f"{index}. {item}" for index, item in enumerate(discrepancies, start=1)),
        "",
        "## Clarifying Questions",
        "",
        "\n".join(f"{index}. {item}" for index, item in enumerate(questions, start=1)),
        "",
    ]
    return "\n".join(sections)


def copy_cell(source, target) -> None:
    target.value = source.value
    if source.has_style:
        target._style = copy.copy(source._style)
    target.font = copy.copy(source.font)
    target.fill = copy.copy(source.fill)
    target.border = copy.copy(source.border)
    target.alignment = copy.copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy.copy(source.protection)


def clear_sheet(ws, max_col: int) -> None:
    for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, 1), min_col=1, max_col=max(max_col, ws.max_column)):
        for cell in row:
            cell.value = None
            cell._style = StyleArray()
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))


def copy_template_layout(source, target, header_rows: int, max_col: int) -> None:
    clear_sheet(target, max_col)
    for row_number in range(1, header_rows + 1):
        for column in range(1, max_col + 1):
            copy_cell(source.cell(row_number, column), target.cell(row_number, column))
        target.row_dimensions[row_number].height = source.row_dimensions[row_number].height
    for merged_range in source.merged_cells.ranges:
        if merged_range.max_row <= header_rows and merged_range.max_col <= max_col:
            target.merge_cells(str(merged_range))
    for column in range(1, max_col + 1):
        letter = openpyxl.utils.get_column_letter(column)
        target.column_dimensions[letter].width = source.column_dimensions[letter].width
        target.column_dimensions[letter].hidden = source.column_dimensions[letter].hidden
    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines


def populate_workbook() -> None:
    workbook = openpyxl.load_workbook(XLSX_PATH)
    field_template = workbook["QDE_BD_Fields"]
    story_template = workbook["QDE_BD_US"]
    field_ws = workbook["QDE_OCCUP_Fields"]
    story_ws = workbook["QDE_OCCUP_US"]

    copy_template_layout(field_template, field_ws, 1, 11)
    for row_number, values in enumerate(fields, start=2):
        for column, value in enumerate(values, start=1):
            target = field_ws.cell(row_number, column, value)
            source = field_template.cell(2, column)
            if source.has_style:
                target._style = copy.copy(source._style)
            target.alignment = copy.copy(source.alignment)
        field_ws.row_dimensions[row_number].height = 84
    field_ws.auto_filter.ref = f"A1:K{len(fields) + 1}"

    copy_template_layout(story_template, story_ws, 2, 14)
    for row_number, story in enumerate(stories, start=3):
        assumptions = "\n".join(f"• {item}" for item in story["assumptions"])
        integrations = "\n".join(f"• {item}" for item in story["integrations"])
        values = [
            story["id"],
            None,
            story["title"],
            story["description"],
            None,
            acceptance_text(story),
            STAGE,
            "Sales Officer",
            assumptions,
            integrations,
            None,
            "Yes",
            "Yes",
            "Yes",
        ]
        for column, value in enumerate(values, start=1):
            target = story_ws.cell(row_number, column, value)
            source = story_template.cell(3, column)
            if source.has_style:
                target._style = copy.copy(source._style)
            target.alignment = copy.copy(source.alignment)
        story_ws.row_dimensions[row_number].height = 330
    story_ws.auto_filter.ref = f"A2:N{len(stories) + 2}"

    temporary = XLSX_PATH.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, XLSX_PATH)


def verify() -> None:
    workbook = openpyxl.load_workbook(XLSX_PATH, data_only=False)
    field_ws = workbook["QDE_OCCUP_Fields"]
    story_ws = workbook["QDE_OCCUP_US"]
    assert field_ws.max_row == len(fields) + 1
    assert story_ws.max_row == len(stories) + 2
    assert [field_ws.cell(row, 1).value for row in range(2, len(fields) + 2)] == [row[0] for row in fields]
    assert [story_ws.cell(row, 1).value for row in range(3, len(stories) + 3)] == [
        story["id"] for story in stories
    ]
    assert all(story_ws.cell(row, 12).value == "Yes" for row in range(3, len(stories) + 3))
    assert all(story_ws.cell(row, 13).value == "Yes" for row in range(3, len(stories) + 3))
    assert all(story_ws.cell(row, 14).value == "Yes" for row in range(3, len(stories) + 3))

    markdown = MD_PATH.read_text(encoding="utf-8")
    assert markdown.count("### US-QDE-OCCUP-CV-") == len(stories)
    assert "Out of Scope" not in markdown
    for story in stories:
        assert story["id"] in markdown
        assert story["id"] not in {"US-QDE-OCCUP-CV-000"}
    for forbidden in ["Application__c.", "Applicant__c.", "Applicant_Employment__c.", "Document_CheckList__c."]:
        story_section = markdown.split("## User Stories", 1)[1].split(
            "## Implementation Discrepancies Requiring Resolution", 1
        )[0]
        assert forbidden not in story_section


MD_PATH.write_text(build_markdown(), encoding="utf-8")
populate_workbook()
verify()
print(f"Created {MD_PATH.name}; populated {len(fields)} fields and {len(stories)} user stories.")
