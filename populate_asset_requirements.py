from __future__ import annotations

import copy
import os
from pathlib import Path

import openpyxl
from openpyxl.styles.cell_style import StyleArray


ROOT = Path(r"C:\Users\mudupi\Desktop\Cursor\Moga Dev")
MD_PATH = ROOT / "CV_DDE_Asset_Details_Requirements.md"
XLSX_PATH = ROOT / "User Stories_CE.xlsx"

STAGE = "DDE"
SUB_STAGE = "Asset Details"
PERSONA = "Sales Officer"


def field(
    name: str,
    obj: str,
    field_type: str,
    values: str | None,
    section: str,
    validation: str,
    error: str | None = None,
    control: str | None = None,
    expression: str | None = None,
) -> list[str | None]:
    return [
        name,
        obj,
        field_type,
        values,
        STAGE,
        SUB_STAGE,
        section,
        validation,
        error,
        control,
        expression,
    ]


states = "Assam; Maharashtra; Karnataka; Delhi; Gujarat"
districts = (
    "Assam: Kamrup, Dibrugarh, Jorhat, Nagaon, Tinsukia; "
    "Maharashtra: Mumbai, Pune, Nagpur, Nashik, Thane; "
    "Karnataka: Bengaluru Urban, Mysuru, Belagavi, Dharwad, Shivamogga; "
    "Delhi: Central Delhi, South Delhi, West Delhi, North Delhi, East Delhi; "
    "Gujarat: Ahmedabad, Surat, Vadodara, Rajkot, Bhavnagar"
)
dealers = "ABC Motors; City Auto Hub; Prime Wheels"
tenures_live = "6; 12; 18; 24; 30; 36; 42; 48; 54; 60; 66; 72; 78; 84"
vehicle_subcategories = (
    "Used - Broker Sale; Used - Open Market; Used - Pre-Owned; Used - Repo Sale"
)


fields: list[list[str | None]] = [
    # Asset list and routing.
    field(
        "Asset / Vehicle ID",
        "Asset__c.Name",
        "Display text",
        None,
        "Asset List",
        "Read-only identifier for the asset card.",
    ),
    field(
        "Asset Type",
        "Asset__c.Vehicle_type__c",
        "Display text",
        "NEW; USED; REFINANCE",
        "Asset List",
        "Read-only on the list and used to route the selected asset to the New or Used/Refinance detail flow.",
    ),
    field(
        "Make",
        "Asset__c.Make_Name__c",
        "Display text",
        None,
        "Asset List",
        "Read-only summary of the selected make.",
    ),
    field(
        "Model",
        "Asset__c.Model_Name__c",
        "Display text",
        None,
        "Asset List",
        "Read-only summary of the selected model.",
    ),
    field(
        "Variant",
        "Asset__c.Variant_Name__c",
        "Display text",
        None,
        "Asset List",
        "Read-only summary of the selected variant.",
    ),
    field(
        "Number of Assets",
        "Application_Line_Item__c.Assets__r",
        "Calculated count",
        None,
        "Asset List",
        "Read-only count of assets under the line item.",
    ),
    field(
        "Asset Status",
        "Application_Line_Item__c.Stage_Completed__c",
        "Calculated status",
        "IN PROGRESS; COMPLETED",
        "Asset List",
        "Read-only status derived from the applicable completion checks.",
    ),
    field(
        "Clone Quantity",
        "Transient UI value",
        "Number",
        None,
        "Asset List — Clone",
        "Required when Clone is selected; whole number of at least 1. Clone is available only for New assets.",
        "Enter a valid quantity of at least 1.",
        "Asset Type",
        "Visible in the Clone dialog when Asset Type = New.",
    ),

    # New vehicle selection and pricing.
    field(
        "Make",
        "Asset__c.Make__c",
        "Searchable master lookup",
        "Product-filtered active make master values",
        "New Asset — Vehicle Selection",
        "Required; a master option must be selected.",
        "Please fill mandatory fields: Make.",
        "Product Type",
        "Options are filtered by the product type resolved for the application.",
    ),
    field(
        "Model",
        "Asset__c.Model__c",
        "Searchable dependent master lookup",
        "Product and make-filtered active model master values",
        "New Asset — Vehicle Selection",
        "Required; changing Make clears Model, Variant, fetched pricing, viability, and Hire vs Buy completion.",
        "Please fill mandatory fields: Model.",
        "Make",
        "Enabled and populated after a Make master value is selected.",
    ),
    field(
        "Variant",
        "Asset__c.Variant__c",
        "Searchable dependent master lookup",
        "Product, make, and model-filtered active variant master values",
        "New Asset — Vehicle Selection",
        "Required; changing Model clears Variant and downstream results.",
        "Please fill mandatory fields: Variant.",
        "Make + Model",
        "Enabled and populated after Make and Model master values are selected.",
    ),
    field(
        "Number of Vehicles",
        "Asset__c.Number_of_Vehicles__c",
        "Number",
        None,
        "New Asset — Vehicle Selection",
        "Required; whole number of at least 1. Disabled until Variant is selected. A value above 1 creates sibling assets.",
        "Please fill mandatory fields: Number of Vehicles.",
        "Variant",
        "Enabled after Variant is selected.",
    ),
    field(
        "1st Year Insurance Premium",
        "Asset__c.X1st_yr_Insurance_Premium__c",
        "Currency",
        None,
        "New Asset — Pricing Details",
        "Optional editable amount on the live screen.",
        None,
        "Fetch Price",
        "Visible after Fetch Price succeeds or persisted pricing exists.",
    ),
    field(
        "Basic Price",
        "Asset__c.Basic_Price__c",
        "Currency",
        None,
        "New Asset — Pricing Details",
        "Required; populated by Fetch Price and editable. Recalculates Ex Showroom Price.",
        "Please fill mandatory fields: Basic Price.",
        "Fetch Price",
        "Visible after Fetch Price succeeds or persisted pricing exists.",
    ),
    field(
        "Discount on Basic Price",
        "Asset__c.Discount_on_Basic_Price__c",
        "Currency",
        None,
        "New Asset — Pricing Details",
        "Required; zero is accepted. Recalculates Ex Showroom Price.",
        "Please fill mandatory fields: Discount on Basic Price.",
        "Fetch Price",
        "Visible after Fetch Price succeeds or persisted pricing exists.",
    ),
    field(
        "GST Amount",
        "Asset__c.GST_Amount__c",
        "Currency",
        None,
        "New Asset — Pricing Details",
        "Required; populated by Fetch Price and editable. Recalculates Ex Showroom Price.",
        "Please fill mandatory fields: GST Amount.",
        "Fetch Price",
        "Visible after Fetch Price succeeds or persisted pricing exists.",
    ),
    field(
        "Other Charges",
        "Asset__c.Other_charges__c",
        "Currency",
        None,
        "New Asset — Pricing Details (live screen)",
        "Optional on the live screen. This deployed field is absent from the reviewed local page and save mapping.",
        None,
        "Fetch Price",
        "Visible after Fetch Price succeeds on the live screen.",
    ),
    field(
        "RTO / Road Tax",
        "Asset__c.RTO_Road_Tax_New__c",
        "Currency",
        None,
        "New Asset — Pricing Details (live screen)",
        "Required on the live screen. This deployed field is absent from the reviewed local page and save mapping.",
        "Please fill mandatory fields: RTO / Road Tax.",
        "Fetch Price",
        "Visible after Fetch Price succeeds on the live screen.",
    ),
    field(
        "Motor Insurance Premium",
        "Asset__c.Motor_Insurance_Premium__c",
        "Currency",
        None,
        "New Asset — Pricing Details (repository)",
        "Optional in the reviewed repository but not displayed on the live screen.",
        None,
        "Fetch Price",
        "Repository-only field requiring alignment with the deployed Other Charges and RTO / Road Tax controls.",
    ),
    field(
        "Ex Showroom Price",
        "Asset__c.Ex_showroom_price__c",
        "Currency",
        None,
        "New Asset — Pricing Details",
        "Required and editable; current calculation is Basic Price minus Discount plus GST.",
        "Please fill mandatory fields: Ex Showroom Price.",
        "Basic Price + Discount + GST",
        "Recalculated when a contributing amount changes.",
    ),
    field(
        "On Road Price",
        "Asset__c.On_Road_price__c",
        "Calculated currency",
        None,
        "New Asset — Pricing Details",
        "Read-only. Uses the successful price response when supplied; otherwise follows Ex Showroom Price.",
        None,
        "Fetch Price / Ex Showroom Price",
        "Displayed after pricing is available.",
    ),

    # New dealer and loan.
    field(
        "State Of Utilization Of Asset",
        "Asset__c.State__c",
        "Picklist",
        states,
        "New Asset — Dealer & Registration",
        "Required. Changing State clears District.",
        "Please fill mandatory fields: State Of Utilization Of Asset.",
        "Pricing Details",
        "Visible after required pricing values are complete.",
    ),
    field(
        "District Of Utilization Of Asset",
        "Asset__c.District__c",
        "Dependent picklist",
        districts,
        "New Asset — Dealer & Registration",
        "Required and disabled until State is selected.",
        "Please fill mandatory fields: District Of Utilization Of Asset.",
        "State Of Utilization Of Asset",
        "Options are restricted to the selected State.",
    ),
    field(
        "Dealer/Sub-Dealer Name",
        "Asset__c.Dealer_Sub_dealer_name__c",
        "Picklist",
        dealers,
        "New Asset — Dealer & Registration",
        "Required.",
        "Please fill mandatory fields: Dealer/Sub-Dealer Name.",
        "Pricing Details",
        "Visible after required pricing values are complete.",
    ),
    field(
        "RC Limit Enabled Dealer",
        "Asset__c.RC_retention_applicable__c",
        "Read-only checkbox on live screen",
        "Yes; No",
        "New Asset — Dealer & Registration",
        "Optional/read-only on the live screen; the repository renders it editable and maps it to RC Retention.",
        None,
        "Dealer selection",
        "Displayed with Dealer & Registration.",
    ),
    field(
        "Invoice In The Name Of",
        "Asset__c.Invoice_in_the_name_of__c",
        "Text",
        None,
        "New Asset — Dealer & Registration",
        "Optional.",
        None,
        "Pricing Details",
        "Visible after required pricing values are complete.",
    ),
    field(
        "Vehicle Registered In The Name Of",
        "Asset__c.Vehicle_registered_in_the_name_of__c",
        "Read-only display",
        "Borrower; Co-borrower",
        "New Asset — Dealer & Registration",
        "Read-only on the live screen with Borrower displayed by default.",
        None,
        "Dealer & Registration",
        "Displayed with Dealer & Registration.",
    ),
    field(
        "Required Loan Amount",
        "Asset__c.Required_Total_Loan_Amount__c",
        "Currency",
        None,
        "New Asset — Loan Information",
        "Required.",
        "Please fill mandatory fields: Required Loan Amount.",
        "Dealer + State + District",
        "Visible after required Dealer & Registration values are complete.",
    ),
    field(
        "Maximum Available Tenure",
        "System-derived value",
        "Read-only number",
        "84 months",
        "New Asset — Loan Information",
        "Read-only maximum displayed as 84.",
        None,
        "Loan Information",
        "Displayed with loan fields.",
    ),
    field(
        "Required Tenure",
        "Asset__c.Required_Tenure__c",
        "Picklist",
        tenures_live,
        "New Asset — Loan Information",
        "Required. The live screen offers six-month increments through 84; the reviewed repository generates every month from 1 through 84.",
        "Please fill mandatory fields: Required Tenure.",
        "Dealer + State + District",
        "Visible after required Dealer & Registration values are complete.",
    ),
    field(
        "Required IRR",
        "Asset__c.Required_CRM_IRR__c",
        "Number / rate",
        None,
        "New Asset — Loan Information",
        "Required. Stored as Currency in metadata although functionally a rate.",
        "Please fill mandatory fields: Required IRR.",
        "Dealer + State + District",
        "Visible after required Dealer & Registration values are complete.",
    ),
    field(
        "Funding Based On",
        "Asset__c.Funding_on_ORP__c / Asset__c.Funding_on_Ex_Showroom__c",
        "Radio",
        "On-road Price; Ex-showroom Price",
        "New Asset — Loan Information",
        "Required; exactly one basis must be selected.",
        "Please fill mandatory fields: Funding Based On.",
        "Dealer + State + District",
        "Visible after required Dealer & Registration values are complete.",
    ),

    # Used vehicle dedupe and details.
    field(
        "Vehicle Registration Format",
        "Asset__c.Registration_Number_Format__c",
        "Radio",
        "New Format (HR55AB1234); Old Format (HR-05-AB-1234)",
        "Used / Refinance — Vehicle Dedupe",
        "Required. The selected format does not currently enforce a matching registration-number pattern.",
        "Vehicle registration format is required.",
        "Asset Type",
        "Visible for Used and Refinance assets.",
    ),
    field(
        "Vehicle Registration Number",
        "Asset__c.Vehicle_Registration_number__c",
        "Text",
        None,
        "Used / Refinance — Vehicle Dedupe",
        "Required. No format-specific regular-expression validation is implemented.",
        "Vehicle registration number is required.",
        "Asset Type",
        "Visible for Used and Refinance assets.",
    ),
    field(
        "Vehicle Sub-Category",
        "Asset__c.Vehicle_SubCategory__c",
        "Picklist",
        vehicle_subcategories,
        "Used / Refinance — Vehicle Dedupe",
        "Required. Live/UI values conflict with the restricted metadata values and can fail on save.",
        "Vehicle sub-category is required.",
        "Asset Type",
        "Visible for Used and Refinance assets.",
    ),
    field(
        "Insurance Expiry Date",
        "Asset__c.Ins_Expiry_date__c",
        "Date",
        None,
        "Used / Refinance — Vehicle Dedupe",
        "Optional and may be refreshed from the vehicle-information service.",
        None,
        "Asset Type",
        "Visible for Used and Refinance assets.",
    ),
    field(
        "Parent Deal Number",
        "Asset__c.Parent_Deal_Number__c",
        "Text",
        None,
        "Used / Refinance — Vehicle Dedupe",
        "Optional.",
        None,
        "Asset Type",
        "Visible for Used and Refinance assets.",
    ),
]


used_vehicle_fields = [
    ("Make", "Asset__c.Make__c", "Searchable master lookup", None, "Required.", "Please fill mandatory fields: Make."),
    ("Model", "Asset__c.Model__c", "Searchable dependent master lookup", None, "Required.", "Please fill mandatory fields: Model."),
    ("Variant", "Asset__c.Variant__c", "Searchable dependent master lookup", None, "Required.", "Please fill mandatory fields: Variant."),
    ("Invoice In The Name Of", "Asset__c.Invoice_in_the_name_of__c", "Text", None, "Required.", "Please fill mandatory fields: Invoice In The Name Of."),
    ("Manufacturing Year & Month", "Asset__c.Manufacturer_Year_Month__c", "Text", None, "Required; no YYYY-MM format validation is implemented.", "Please fill mandatory fields: Manufacturing Year & Month."),
    ("Last Owner Name", "Asset__c.Last_owner_name__c", "Text", None, "Required.", "Please fill mandatory fields: Last Owner Name."),
    ("Engine Number", "Asset__c.Engine_number__c", "Text", None, "Required.", "Please fill mandatory fields: Engine Number."),
    ("Chassis Number", "Asset__c.Chassis_number__c", "Text", None, "Required.", "Please fill mandatory fields: Chassis Number."),
    ("Number of Ownerships", "Asset__c.Number_of_ownerships__c", "Picklist", "1; 2; 3; >3", "Required.", "Please fill mandatory fields: Number of Ownerships."),
    ("RC Retention Applicable", "Asset__c.RC_retention_applicable__c", "Picklist", "Yes; No", "Required.", "Please fill mandatory fields: RC Retention Applicable."),
]
for name, obj, field_type, values, validation, error in used_vehicle_fields:
    control = "Get Vehicle Details"
    expression = "Visible after Vehicle Dedupe and the full vehicle-information sequence succeed."
    if name == "Model":
        control = "Get Vehicle Details + Make"
    elif name == "Variant":
        control = "Get Vehicle Details + Make + Model"
    fields.append(
        field(
            name,
            obj,
            field_type,
            values,
            "Used / Refinance — Vehicle Details",
            validation,
            error,
            control,
            expression,
        )
    )


used_shared_fields = [
    ("State Of Utilization Of Asset", "Asset__c.State__c", "Picklist", states, "Required; changing State clears District.", "Please fill mandatory fields: State Of Utilization Of Asset.", "Vehicle Details", "Visible after all required Used Vehicle Details are complete."),
    ("District Of Utilization Of Asset", "Asset__c.District__c", "Dependent picklist", districts, "Required and disabled until State is selected.", "Please fill mandatory fields: District Of Utilization Of Asset.", "State Of Utilization Of Asset", "Options are restricted to the selected State."),
    ("Dealer/Sub-Dealer Name", "Asset__c.Dealer_Sub_dealer_name__c", "Picklist", dealers, "Required.", "Please fill mandatory fields: Dealer/Sub-Dealer Name.", "Vehicle Details", "Visible after all required Used Vehicle Details are complete."),
    ("RC Limit Enabled Dealer", "Asset__c.RC_retention_applicable__c", "Checkbox", "Yes; No", "Optional; writes the same stored value as RC Retention Applicable.", None, "Vehicle Details", "Visible after all required Used Vehicle Details are complete."),
    ("Required Loan Amount", "Asset__c.Required_Total_Loan_Amount__c", "Currency", None, "Required.", "Please fill mandatory fields: Required Loan Amount.", "Dealer + State + District", "Visible after required dealer values are complete."),
    ("Maximum Available Tenure", "System-derived value", "Read-only number", "84 months", "Read-only maximum displayed as 84.", None, "Loan Information", "Displayed with loan fields."),
    ("Required Tenure", "Asset__c.Required_Tenure__c", "Picklist", tenures_live, "Required; live values are six-month increments through 84.", "Please fill mandatory fields: Required Tenure.", "Dealer + State + District", "Visible after required dealer values are complete."),
    ("Required IRR", "Asset__c.Required_CRM_IRR__c", "Number / rate", None, "Required.", "Please fill mandatory fields: Required IRR.", "Dealer + State + District", "Visible after required dealer values are complete."),
    ("Funding Based On", "Asset__c.Funding_on_ORP__c / Asset__c.Funding_on_Ex_Showroom__c", "Radio", "On-road Price; Ex-showroom Price", "Required; exactly one basis must be selected.", "Please fill mandatory fields: Funding Based On.", "Dealer + State + District", "Visible after required dealer values are complete."),
]
for name, obj, field_type, values, validation, error, control, expression in used_shared_fields:
    section = "Used / Refinance — Dealer & Registration"
    if name in {"Required Loan Amount", "Maximum Available Tenure", "Required Tenure", "Required IRR", "Funding Based On"}:
        section = "Used / Refinance — Loan Information"
    fields.append(field(name, obj, field_type, values, section, validation, error, control, expression))


# Viability fields.
fields.extend(
    [
        field(
            "Calculation Basis",
            "Transient UI value",
            "Picklist",
            "Hourly; KMS",
            "Viability — Method Selection",
            "Required. Determines which income and expense branch is displayed. The selection is inferred rather than explicitly persisted.",
            "Select a Calculation Basis.",
            "Viability Check",
            "Visible after loan information is complete and Viability Check is opened.",
        ),
        field(
            "Asset Deployment",
            "Viability_Check__c.Asset_Deployment__c",
            "Picklist",
            "Earth Moving; Material Handling; Road Construction",
            "Viability — Hourly Income",
            "Required.",
            "Please complete all mandatory Hourly viability fields.",
            "Calculation Basis",
            "Visible when Calculation Basis = Hourly.",
        ),
        field(
            "No of Days the Machine Will Work per Month",
            "Viability_Check__c.NO_OF_DAYS_THE_MACHINE_WILL_WORK_PER_MON__c",
            "Number",
            None,
            "Viability — Hourly Income",
            "Required numeric value.",
            "Please complete all mandatory Hourly viability fields.",
            "Calculation Basis",
            "Visible when Calculation Basis = Hourly.",
        ),
        field(
            "No of Hours the Machine Will Work per Day on Average",
            "Viability_Check__c.NO_OF_HOURS_MACHINE_WILL_WORK_PER_DAY__c",
            "Number",
            None,
            "Viability — Hourly Income",
            "Required numeric value.",
            "Please complete all mandatory Hourly viability fields.",
            "Calculation Basis",
            "Visible when Calculation Basis = Hourly.",
        ),
        field(
            "Rate per Hour Estimated to Realise",
            "Viability_Check__c.RATE_PER_HOUR_ESTIMATED_TO_REALISE_RS__c",
            "Currency",
            None,
            "Viability — Hourly Income",
            "Required numeric value.",
            "Please complete all mandatory Hourly viability fields.",
            "Calculation Basis",
            "Visible when Calculation Basis = Hourly.",
        ),
        field(
            "Estimated Average Monthly Income / Cost",
            "Viability_Check__c.ESTIMATED_AVERAGE_INCOME_MONTH__c",
            "Calculated currency",
            None,
            "Viability — Hourly Income",
            "Read-only calculation from days, hours, and rate.",
            None,
            "Hourly income inputs",
            "Visible when Calculation Basis = Hourly.",
        ),
    ]
)


hourly_expenses = [
    ("Fuel Cost per Litre", "Viability_Check__c.FUEL_COST_PER_LITRE_RS__c", "Currency"),
    ("Fuel Consumption per Hour", "Viability_Check__c.FUEL_CONSUMPTION_PER_HOUR__c", "Number"),
    ("Cost of Fuel per Month", "Viability_Check__c.Cost_Of_Fuel_Per_Month__c", "Currency"),
    ("Operator Salary / Month", "Viability_Check__c.OPERATORS_SALARY__c", "Currency"),
    ("Maintenance Parts / Month", "Viability_Check__c.MAINTAINENCE_PARTS_COST__c", "Currency"),
    ("Tyre Cost", "Viability_Check__c.TYRE_COST__c", "Currency"),
    ("Tax / Cess", "Viability_Check__c.MOTOR_TAX_PERMIT_COST__c", "Currency"),
    ("Others", "Viability_Check__c.OTHERS__c", "Currency"),
    ("Total Expenses", "Viability_Check__c.TOTAL_EXPENSES__c", "Calculated currency"),
    ("Gross Surplus", "Viability_Check__c.GROSS_SURPLUS__c", "Calculated currency"),
    ("Instalment Amount", "Viability_Check__c.INSTALMENT_AMOUNT__c", "Currency"),
    ("Net Surplus", "Viability_Check__c.NET_SURPLUS__c", "Calculated currency"),
    ("Viability Ratio", "Viability_Check__c.Viability_Ratio__c", "Calculated number"),
]
for name, obj, field_type in hourly_expenses:
    calculated = field_type.startswith("Calculated")
    fields.append(
        field(
            name,
            obj,
            field_type,
            None,
            "Viability — Hourly Expenses & Result",
            "Read-only calculated value." if calculated else "Required numeric value.",
            None if calculated else "Please complete all mandatory Hourly viability fields.",
            "Calculation Basis",
            "Visible when Calculation Basis = Hourly.",
        )
    )


fields.extend(
    [
        field("Vehicle Route", "Viability_Check__c.Vehicle_Route__c", "Picklist", "Inter State; Within State; Local", "Viability — KMS Income", "Required.", "Please complete all mandatory KMS viability fields.", "Calculation Basis", "Visible when Calculation Basis = KMS."),
        field("From", "Viability_Check__c.From_Location__c", "Text", None, "Viability — KMS Income", "Required.", "Please complete all mandatory KMS viability fields.", "Calculation Basis", "Visible when Calculation Basis = KMS."),
        field("To", "Viability_Check__c.To_Location__c", "Text", None, "Viability — KMS Income", "Required.", "Please complete all mandatory KMS viability fields.", "Calculation Basis", "Visible when Calculation Basis = KMS."),
        field("Trips per Month", "Viability_Check__c.Trips_Per_Month__c", "Number", None, "Viability — KMS Income", "Required whole-number input.", "Please complete all mandatory KMS viability fields.", "Calculation Basis", "Visible when Calculation Basis = KMS."),
        field("No of Kilometers per Trip", "Viability_Check__c.No_Of_Kms_Per_Trip__c", "Number", None, "Viability — KMS Income", "Required numeric value.", "Please complete all mandatory KMS viability fields.", "Calculation Basis", "Visible when Calculation Basis = KMS."),
        field("Average Rate per Trip", "Viability_Check__c.Average_Rate_Per_Trip__c", "Currency", None, "Viability — KMS Income", "Required numeric value.", "Please complete all mandatory KMS viability fields.", "Calculation Basis", "Visible when Calculation Basis = KMS."),
        field("Estimated Average Income per Month", "Viability_Check__c.ESTIMATED_AVERAGE_INCOME_MONTH__c", "Calculated currency", None, "Viability — KMS Income", "Read-only calculation from trips and average rate.", None, "KMS income inputs", "Visible when Calculation Basis = KMS."),
    ]
)


kms_expenses = [
    ("Fuel Rate per Liter", "Viability_Check__c.FUEL_COST_PER_LITRE_RS__c", "Currency"),
    ("Cost of Fuel per Trip", "Viability_Check__c.Cost_Of_Fuel_Per_Trip__c", "Currency"),
    ("Cost of Fuel per Month", "Viability_Check__c.Cost_Of_Fuel_Per_Month__c", "Currency"),
    ("Salaries", "Viability_Check__c.OPERATORS_SALARY__c", "Currency"),
    ("Maintenance", "Viability_Check__c.MAINTAINENCE_PARTS_COST__c", "Currency"),
    ("Tyre Cost", "Viability_Check__c.TYRE_COST__c", "Currency"),
    ("Motor Tax / Permit Cost / Cess", "Viability_Check__c.MOTOR_TAX_PERMIT_COST__c", "Currency"),
    ("Others", "Viability_Check__c.OTHERS__c", "Currency"),
    ("Total Expenses", "Viability_Check__c.TOTAL_EXPENSES__c", "Calculated currency"),
    ("Gross Surplus", "Viability_Check__c.GROSS_SURPLUS__c", "Calculated currency"),
    ("Instalment Amount", "Viability_Check__c.INSTALMENT_AMOUNT__c", "Currency"),
    ("Net Surplus", "Viability_Check__c.NET_SURPLUS__c", "Calculated currency"),
    ("Viability Ratio", "Viability_Check__c.Viability_Ratio__c", "Calculated number"),
]
for name, obj, field_type in kms_expenses:
    calculated = field_type.startswith("Calculated")
    fields.append(
        field(
            name,
            obj,
            field_type,
            None,
            "Viability — KMS Expenses & Result",
            "Read-only calculated value." if calculated else "Required numeric value.",
            None if calculated else "Please complete all mandatory KMS viability fields.",
            "Calculation Basis",
            "Visible when Calculation Basis = KMS.",
        )
    )


# Hire versus Buy fields.
hire_fields = [
    ("Vehicle Model", "Viability_Check__c.Vehicle_Model__c", "Read-only text", "Read-only value from the selected asset.", None),
    ("Type of Work", "Viability_Check__c.Type_of_work__c", "Text", "Optional; Save currently permits blank.", None),
    ("No of Hours / Day", "Viability_Check__c.NO_OF_HOURS_MACHINE_WILL_WORK_PER_DAY__c", "Number", "Optional; when entered, must be less than 24.", "No of Hours / Day must be less than 24."),
    ("Working Days / Month", "Viability_Check__c.NO_OF_DAYS_THE_MACHINE_WILL_WORK_PER_MON__c", "Number", "Optional; when entered, must be 31 or less.", "Working Days / Month must be less than or equal to 31."),
    ("Monthly Running in Hours", "Viability_Check__c.Monthly_running_in_hours__c", "Calculated number", "Read-only: Hours per Day multiplied by Working Days per Month.", None),
    ("Rate / Hour", "Viability_Check__c.RATE_PER_HOUR_ESTIMATED_TO_REALISE_RS__c", "Currency", "Optional; blank and negative values are not currently blocked.", None),
    ("Monthly Charges for Hire", "Viability_Check__c.Monthly_charges_for_hire__c", "Calculated currency", "Read-only: Monthly Running Hours multiplied by Rate per Hour.", None),
    ("Fuel Average / Hour", "Viability_Check__c.FUEL_CONSUMPTION_PER_HOUR__c", "Number", "Optional; blank and negative values are not currently blocked.", None),
    ("Cost per Litre", "Viability_Check__c.FUEL_COST_PER_LITRE_RS__c", "Currency", "Optional; blank and negative values are not currently blocked.", None),
    ("Fuel Cost", "Viability_Check__c.Cost_Of_Fuel_Per_Month__c", "Calculated currency", "Read-only calculated operating cost.", None),
    ("Drivers Salary & Allowances", "Viability_Check__c.OPERATORS_SALARY__c", "Currency", "Optional; blank and negative values are not currently blocked.", None),
    ("Insurance / Taxes etc", "Viability_Check__c.MOTOR_TAX_PERMIT_COST__c", "Currency", "Optional; blank and negative values are not currently blocked.", None),
    ("Repair / Servicing / Misc", "Viability_Check__c.MAINTAINENCE_PARTS_COST__c", "Currency", "Optional; blank and negative values are not currently blocked.", None),
    ("EMI", "Viability_Check__c.INSTALMENT_AMOUNT__c", "Currency", "Optional; blank and negative values are not currently blocked.", None),
    ("Existing EMI", "Viability_Check__c.Emi__c", "Currency", "Optional; blank and negative values are not currently blocked.", None),
    ("Total EMI", "Viability_Check__c.Total_EMI__c", "Calculated currency", "Read-only: EMI plus Existing EMI.", None),
    ("Total Expenses", "Viability_Check__c.TOTAL_EXPENSES__c", "Calculated currency", "Read-only total ownership cost.", None),
    ("Hire vs Buy Difference", "Viability_Check__c.HireVsBuy_Amt__c", "Calculated currency", "Read-only: Hire Charges minus ownership expenses.", None),
    ("Hire vs Buy Is Favourable", "Viability_Check__c.HireVsBuy_Decision__c", "Calculated text", "FAVOURABLE only when the difference is greater than zero; otherwise NON FAVOURABLE.", None),
]
for name, obj, field_type, validation, error in hire_fields:
    fields.append(
        field(
            name,
            obj,
            field_type,
            "FAVOURABLE; NON FAVOURABLE" if name == "Hire vs Buy Is Favourable" else None,
            "Hire vs Buy",
            validation,
            error,
            "Product Type + Loan Information",
            "Available for SCV, LCV, MHCV, and CE after loan information is complete.",
        )
    )


def story(
    sid: str,
    title: str,
    description: str,
    scenarios: list[tuple[str, list[str]]],
    assumptions: list[str],
    integrations: list[str],
    scope: list[str],
) -> dict:
    return {
        "id": sid,
        "title": title,
        "description": description,
        "scenarios": scenarios,
        "assumptions": assumptions,
        "integrations": integrations,
        "scope": scope,
    }


stories = [
    story(
        "US-DDE-ASSET-CV-001",
        "Manage Assets for the Application",
        "As a Sales Officer, I want to add, review, edit, clone, and remove application assets so that every financed vehicle is represented correctly.",
        [
            ("Positive Scenario — Add asset", ["Given the application is at DDE Asset Details", "When Add Asset is selected", "Then a new asset and line item are created and the applicable detail flow opens"]),
            ("Positive Scenario — Edit and review", ["Given an asset exists", "When Edit or View is selected", "Then the correct detail flow opens with the saved values"]),
            ("Positive Scenario — Clone", ["Given a New asset exists", "When a valid clone quantity is confirmed", "Then separate New asset line items are created using the source identification details"]),
            ("Validation Scenario", ["Given clone quantity is blank, zero, negative, or not a whole number", "When Clone is confirmed", "Then cloning is blocked with a clear validation message"]),
            ("Deletion Scenario", ["Given an asset is deleted", "When deletion succeeds", "Then its card is removed and the parent line item is removed only when no sibling asset remains"]),
            ("Edge Case Scenario", ["Given the application has no assets", "When Next is selected", "Then navigation is blocked until at least one complete asset exists"]),
        ],
        ["The current application permits multiple asset line items.", "Clone applies only to New assets."],
        ["Asset management uses internal record services; no external integration is triggered by Add, Clone, Edit, View, or Delete."],
        ["Asset list, counts, status, pagination, Add, Edit, View, Clone, and Delete."],
    ),
    story(
        "US-DDE-ASSET-CV-002",
        "Route New, Used, and Refinance Asset Profiles",
        "As a Sales Officer, I want each asset to open the correct profile so that the relevant identification and assessment steps are applied.",
        [
            ("Positive Scenario — New", ["Given the asset type is New", "When the asset is opened", "Then Vehicle Selection and price retrieval are displayed"]),
            ("Positive Scenario — Used or Refinance", ["Given the asset type is Used or Refinance", "When the asset is opened", "Then vehicle dedupe and registration-based retrieval are displayed"]),
            ("Navigation Scenario", ["Given a New asset is complete", "When the stage advances", "Then the next DDE step is Loan Terms", "And Used or Refinance routes through Valuation before Loan Terms"]),
            ("Edge Case Scenario", ["Given the asset type is blank or unsupported", "When routing is evaluated", "Then the system must not silently treat it as New without an approved business rule"]),
        ],
        ["Asset type is established before DDE Asset Details.", "Mixed asset types require an application-level routing decision."],
        ["No external integration is triggered by profile routing.", "The journey-navigation service resolves the next DDE step."],
        ["New, Used, and Refinance profile selection and route determination."],
    ),
    story(
        "US-DDE-ASSET-CV-003",
        "Select Vehicle Make, Model, and Variant",
        "As a Sales Officer, I want to select a valid product-specific vehicle hierarchy so that pricing and policy checks use an approved vehicle configuration.",
        [
            ("Positive Scenario", ["Given a New asset is being captured", "When Make, Model, and Variant are selected in sequence", "Then only active product-compatible master values are accepted"]),
            ("Dependency Scenario", ["Given Make changes", "When the change is accepted", "Then Model and Variant are cleared", "And changing Model clears Variant"]),
            ("Reset Scenario", ["Given price, viability, or Hire vs Buy has already been completed", "When Make, Model, Variant, or vehicle count changes", "Then affected downstream results are invalidated and must be rerun"]),
            ("Negative Scenario", ["Given the dependent master contains no valid Model or Variant", "When the Sales Officer searches", "Then the absence of results is shown and Fetch Price remains blocked"]),
            ("Validation Scenario", ["Given a typed value was not selected from the master", "When the Sales Officer continues", "Then the value is not treated as a valid selection"]),
        ],
        ["The master-data hierarchy and product mapping are maintained before application capture."],
        ["An internal vehicle master lookup supplies Make, Model, and Variant.", "There is no manual free-text fallback when a master value is missing."],
        ["Vehicle hierarchy search, dependencies, selection, reset behavior, and unavailable-master handling."],
    ),
    story(
        "US-DDE-ASSET-CV-004",
        "Fetch and Review New-Vehicle Pricing",
        "As a Sales Officer, I want to retrieve and review vehicle pricing so that the asset cost used for financing is transparent and accurate.",
        [
            ("Positive Scenario", ["Given Make, Model, Variant, and a valid vehicle count are complete", "When Fetch Price is selected", "Then pricing is returned and the Pricing Details section is displayed"]),
            ("Calculation Scenario", ["Given Basic Price, Discount, or GST changes", "When the value is accepted", "Then Ex Showroom Price is recalculated as Basic Price minus Discount plus GST"]),
            ("Validation Scenario", ["Given a required pricing value is missing", "When the Sales Officer attempts to continue", "Then completion is blocked and the missing field is identified"]),
            ("Failure Scenario", ["Given the pricing request fails or times out", "When the response is received", "Then an error is shown, pricing completion is not recorded, and Fetch Price can be retried"]),
            ("Persistence Scenario", ["Given pricing is saved", "When the asset is reopened", "Then the saved values are restored without requiring another fetch unless a vehicle selection input changes"]),
        ],
        ["The final approved pricing field set must reconcile the live screen with the repository.", "Currency values are expressed in Indian Rupees."],
        ["The current price service returns simulated fixed values rather than a production quotation.", "Failure blocks the dependent pricing, dealer, and loan sections; no pre-success manual-entry route exists."],
        ["Price retrieval, editable pricing components, calculations, validation, failure, retry, and persistence."],
    ),
    story(
        "US-DDE-ASSET-CV-005",
        "Capture Dealer, Utilization, and Registration Details",
        "As a Sales Officer, I want to capture dealer and vehicle-utilization details so that the financed asset is associated with the correct supplier and operating location.",
        [
            ("Positive Scenario", ["Given required pricing is complete", "When Dealer & Registration is displayed", "Then State, District, and Dealer are required"]),
            ("Dependency Scenario", ["Given a State is selected", "When the District list opens", "Then only districts for that State are available", "And changing State clears the previous District"]),
            ("Validation Scenario", ["Given State, District, or Dealer is missing", "When the Sales Officer continues", "Then completion is blocked and the missing value is identified"]),
            ("Review Scenario", ["Given RC Limit and registration-name values are displayed", "When the asset is reviewed", "Then read-only and editable behavior follows the approved dealer policy"]),
        ],
        ["Dealer, State, and District lists are approved reference data.", "Vehicle Registered In The Name Of defaults to Borrower unless business rules specify otherwise."],
        ["The current lists are internal hardcoded values; no external dealer or supplier lookup is invoked."],
        ["Dealer selection, utilization State and District dependency, invoice name, RC indicator, and registration name."],
    ),
    story(
        "US-DDE-ASSET-CV-006",
        "Capture Requested Asset Finance Terms",
        "As a Sales Officer, I want to capture the requested loan amount, tenure, rate, and funding basis so that viability uses the intended financing request.",
        [
            ("Positive Scenario", ["Given dealer details are complete", "When Loan Information is displayed", "Then Loan Amount, Tenure, IRR, and Funding Basis are required"]),
            ("Validation Scenario", ["Given a required loan value is missing", "When Viability or Next is selected", "Then the action is blocked and the missing value is identified"]),
            ("Dependency Scenario", ["Given Funding Based On is selected", "When the value is saved", "Then exactly one of On-road Price or Ex-showroom Price is recorded"]),
            ("Boundary Scenario", ["Given tenure is selected", "When the list opens", "Then the live journey offers approved six-month increments up to the displayed maximum of 84 months"]),
            ("Edge Case Scenario", ["Given IRR is entered", "When saved", "Then it is treated as a rate and not as a currency amount"]),
        ],
        ["Maximum Available Tenure is supplied by policy and is currently displayed as 84 months."],
        ["No external finance-quotation service is invoked by these fields.", "The values become prerequisites for viability and policy assessment."],
        ["Requested loan amount, tenure, IRR, funding basis, bounds, and gating."],
    ),
    story(
        "US-DDE-ASSET-CV-007",
        "Assess Asset Viability Using Hourly Operations",
        "As a Sales Officer, I want to assess viability from hourly machine utilization so that expected income, expenses, instalment burden, and surplus are evaluated.",
        [
            ("Positive Scenario", ["Given Calculation Basis is Hourly", "When deployment, working days, hours, rate, and expense inputs are completed", "Then monthly income, expenses, surplus, and viability ratio are calculated"]),
            ("Validation Scenario", ["Given any required Hourly value is missing or invalid", "When Viability Check is selected", "Then the assessment is not saved and the missing or invalid value is identified"]),
            ("Integration Scenario", ["Given complete Hourly inputs", "When Viability Check is selected", "Then the viability assessment is processed and saved only after a successful response"]),
            ("Failure and Retry Scenario", ["Given Hourly viability processing fails", "When the failure is shown", "Then the assessment remains incomplete and Retry reruns the assessment"]),
            ("Persistence Scenario", ["Given a viability result is saved", "When the asset is reopened", "Then the saved inputs and calculated result are available for review"]),
        ],
        ["All mandatory income and expense values are available to the Sales Officer.", "The approved ratio threshold is managed outside this capture story unless specified."],
        ["Hourly viability currently uses a simulated processing response followed by internal persistence.", "Failure blocks New-asset completion."],
        ["Hourly income, expenses, calculations, save, failure, retry, and review."],
    ),
    story(
        "US-DDE-ASSET-CV-008",
        "Assess Asset Viability Using Route and Kilometres",
        "As a Sales Officer, I want to assess viability from route, trips, and distance so that transport income and operating costs are evaluated.",
        [
            ("Positive Scenario", ["Given Calculation Basis is KMS", "When Route, From, To, Trips, Kilometres, Rate, and expenses are completed", "Then income, expenses, surplus, and viability ratio are calculated"]),
            ("Dependency Scenario", ["Given Calculation Basis changes from Hourly to KMS", "When the KMS branch is displayed", "Then only KMS-specific inputs are required for the active assessment"]),
            ("Validation Scenario", ["Given a required KMS value is missing or invalid", "When Viability Check is selected", "Then saving is blocked and the missing or invalid value is identified"]),
            ("Persistence Scenario", ["Given a KMS result is saved", "When the asset is revisited", "Then the route, expense, and result values are restored"]),
            ("Edge Case Scenario", ["Given values exist in both branches", "When the screen is reloaded", "Then the system must restore the approved Calculation Basis rather than infer an ambiguous branch"]),
        ],
        ["Route and trip inputs represent average monthly operations."],
        ["KMS viability is calculated and saved internally; no external service is currently invoked."],
        ["KMS route, trips, distance, income, expenses, calculations, persistence, and branch switching."],
    ),
    story(
        "US-DDE-ASSET-CV-009",
        "Compare Hiring Cost with Asset Ownership",
        "As a Sales Officer, I want to compare hiring charges with ownership expenses so that the financial preference for buying the asset is visible.",
        [
            ("Positive Scenario", ["Given an applicable product and complete loan information", "When Hire vs Buy inputs are entered", "Then running hours, hire charges, fuel cost, EMI, total expenses, difference, and favourability are calculated"]),
            ("Calculation Scenario", ["Given Hire Charges exceed ownership expenses", "When the comparison is calculated", "Then the result is FAVOURABLE", "And a zero or negative difference is NON FAVOURABLE"]),
            ("Validation Scenario", ["Given Hours per Day is 24 or more or Working Days exceeds 31", "When Save is selected", "Then saving is blocked with the applicable limit message"]),
            ("Negative Scenario", ["Given required business inputs are blank or negative", "When Save is selected", "Then the final requirement must determine whether completion is blocked; the current implementation permits most such values"]),
            ("Persistence Scenario", ["Given the comparison is saved", "When the asset is reopened", "Then the saved comparison opens in review mode"]),
        ],
        ["Hire vs Buy is mandatory for New SCV, LCV, MHCV, and CE products in the current flow.", "Business-required inputs and non-negative rules require confirmation."],
        ["No external Hire vs Buy service is invoked; calculations and persistence are internal."],
        ["Hiring inputs, ownership inputs, calculations, validation, save, and completed review."],
    ),
    story(
        "US-DDE-ASSET-CV-010",
        "Run Used-Vehicle Dedupe",
        "As a Sales Officer, I want to check a used or refinance vehicle for existing deals so that duplicate or restricted vehicles do not proceed.",
        [
            ("Positive Scenario", ["Given registration format, sub-category, and registration number are complete", "When Vehicle Dedupe is selected", "Then the result permits continuation only when the proceed indicator is positive"]),
            ("Validation Scenario", ["Given a mandatory dedupe input is missing", "When Vehicle Dedupe is selected", "Then the request is not sent and the missing input is identified"]),
            ("Negative Scenario", ["Given dedupe returns a non-proceed result", "When the response is received", "Then Get Vehicle Details remains blocked and the response reason is shown"]),
            ("Failure and Retry Scenario", ["Given dedupe fails or times out", "When the error is displayed", "Then the Sales Officer can Retry and Next remains blocked"]),
            ("Reset Scenario", ["Given dedupe previously succeeded", "When any dedupe request input changes", "Then the previous success must be invalidated before further retrieval"]),
        ],
        ["The authoritative registration format and sub-category lists are approved.", "A positive dedupe decision is mandatory before vehicle-detail retrieval."],
        ["The current dedupe response is simulated and always returns a positive proceed indicator.", "No manual bypass exists."],
        ["Dedupe inputs, trigger, decision, failure, retry, reset, and gating."],
    ),
    story(
        "US-DDE-ASSET-CV-011",
        "Retrieve Used-Vehicle Registration and Master Details",
        "As a Sales Officer, I want to retrieve registration and vehicle-master information so that used-vehicle details can be reviewed and completed accurately.",
        [
            ("Positive Scenario", ["Given dedupe succeeds", "When Get Vehicle Details is selected", "Then registration, master, and final vehicle-detail steps run in sequence"]),
            ("Data Mapping Scenario", ["Given the sequence succeeds", "When the response is applied", "Then owner, engine, chassis, manufacture date, insurance date, ownership count, Make, Model, and Variant are populated"]),
            ("Failure Scenario", ["Given any step fails or times out", "When the failure is shown", "Then later steps do not run and Vehicle Details remains unavailable"]),
            ("Retry Scenario", ["Given a retrieval step fails", "When Retry is selected", "Then the full sequence restarts from the first step"]),
            ("Manual Entry Scenario", ["Given upstream data is unavailable or unmapped", "When the failure persists", "Then manual entry remains blocked until an approved fallback is defined"]),
        ],
        ["The returned Make, Model, and Variant map to valid active master values.", "The Sales Officer reviews fetched information before completion."],
        ["The current registration report, vehicle master, and completion responses are simulated with fixed values.", "Failure blocks downstream vehicle details and has no manual bypass."],
        ["Sequential retrieval, data mapping, progress, failure, retry, and fallback behavior."],
    ),
    story(
        "US-DDE-ASSET-CV-012",
        "Complete Used-Vehicle Identification and Valuation Inputs",
        "As a Sales Officer, I want to review and complete used-vehicle identification details so that the asset can proceed to valuation and loan assessment.",
        [
            ("Positive Scenario", ["Given vehicle retrieval succeeds", "When the Used Vehicle Details section opens", "Then Make, Model, Variant, invoice name, manufacture period, owner, engine, chassis, ownership count, and RC retention are required"]),
            ("Validation Scenario", ["Given any required vehicle field is missing", "When the Sales Officer continues", "Then completion is blocked and the missing value is identified"]),
            ("Format Scenario", ["Given Registration Number or Manufacturing Year and Month is entered", "When the value is validated", "Then it must follow the approved format rather than accepting arbitrary text"]),
            ("Navigation Scenario", ["Given Used or Refinance details, dealer, loan, and viability are complete", "When the asset stage advances", "Then the journey proceeds to Valuation"]),
            ("Persistence Scenario", ["Given valid details are saved", "When the asset is reopened", "Then all reviewed and entered values are restored"]),
        ],
        ["Valuation is a subsequent DDE step for Used and Refinance assets.", "RC Retention and RC Limit require separate business definitions if both are retained."],
        ["Fetched vehicle data originates from the registration and master sequence.", "Final save and navigation are internal."],
        ["Used identification, validation, dealer and loan dependencies, save, persistence, and valuation route."],
    ),
    story(
        "US-DDE-ASSET-CV-013",
        "Persist and Synchronize Multi-Vehicle Asset Details",
        "As a Sales Officer, I want multi-vehicle details to remain consistent so that every sibling asset carries the approved pricing, dealer, and finance information.",
        [
            ("Positive Scenario", ["Given Number of Vehicles is greater than 1", "When Fetch Price runs", "Then the required sibling assets are created under the same line item"]),
            ("Synchronization Scenario", ["Given the primary New asset is saved", "When sibling assets exist", "Then approved vehicle, pricing, dealer, and loan values are synchronized to the siblings"]),
            ("Edge Case Scenario — Reduce count", ["Given sibling assets already exist", "When Number of Vehicles is reduced", "Then the system must not leave unintended extra assets without confirmation"]),
            ("Clone Scenario", ["Given a New asset is cloned", "When cloning completes", "Then the clone contains identification details only unless broader copy behavior is approved"]),
            ("Persistence Scenario", ["Given the list is refreshed", "When the line item is reopened", "Then sibling count and saved values remain consistent"]),
        ],
        ["Sibling assets represent identical vehicles under one line item unless individually edited later."],
        ["No external integration is triggered by synchronization.", "Pricing and viability results are not independently fetched for each sibling in the current flow."],
        ["Sibling creation, synchronization, reduction behavior, clone scope, and persistence."],
    ),
    story(
        "US-DDE-ASSET-CV-014",
        "Handle Asset Integrations and Recovery",
        "As a Sales Officer, I want clear progress and recovery behavior for every Asset Details integration so that failures do not create false completion.",
        [
            ("Integration Scenario", ["Given an integration is triggered", "When it is processing", "Then the current step and progress are shown and duplicate submission is prevented"]),
            ("Failure Scenario", ["Given a request fails, times out, or returns an unsuccessful business response", "When the response is handled", "Then completion is not recorded and downstream actions remain blocked"]),
            ("Retry Scenario", ["Given retry is allowed", "When Retry is selected", "Then the approved request sequence is rerun without duplicating completed asset records"]),
            ("Data Change Scenario", ["Given integration inputs change after success", "When the change is accepted", "Then the affected success state is cleared and the integration must be rerun"]),
            ("Manual Fallback Scenario", ["Given repeated upstream failure", "When no manual path is configured", "Then the Sales Officer is informed that the journey is blocked and escalation guidance is provided"]),
        ],
        ["Production service contracts, timeout periods, retry limits, and support ownership will be approved before go-live."],
        ["Covers vehicle master lookup, Fetch Price, vehicle dedupe, registration report, deal master, vehicle details, and Hourly viability.", "Current service bridges are partly or wholly simulated."],
        ["Progress, success, business failure, technical failure, timeout, retry, input reset, gating, and manual fallback."],
    ),
    story(
        "US-DDE-ASSET-CV-015",
        "Complete, Navigate, and Review Asset Details",
        "As a Sales Officer, I want Asset Details to complete only after every applicable requirement succeeds so that the application advances with accurate and reviewable asset information.",
        [
            ("Completion Scenario — New", ["Given required New-asset fields, pricing, dealer, loan, viability, and applicable Hire vs Buy are saved", "When the asset returns to the list", "Then the asset is marked complete"]),
            ("Completion Scenario — Used or Refinance", ["Given dedupe, retrieval, required vehicle details, dealer, loan, and viability are saved", "When completion is evaluated", "Then the full required form determines status"]),
            ("Navigation Scenario", ["Given every displayed asset is complete", "When Next is selected", "Then New routes to Loan Terms and Used or Refinance routes to Valuation"]),
            ("Negative Scenario", ["Given any asset is incomplete or an integration is unresolved", "When Next is selected", "Then navigation is blocked with the incomplete requirement identified"]),
            ("Review Scenario", ["Given Asset Details is completed and the application moves forward", "When the stage is revisited", "Then saved fields and completed analyses are displayed in read-only mode"]),
            ("Back Scenario", ["Given unsaved changes exist", "When Back to Line Items is selected", "Then the user is not led to believe the unsaved values were persisted"]),
        ],
        ["All assets must be complete before application-level navigation.", "Asset-specific document upload is not currently present and any required documents need an approved location and gate."],
        ["The journey-navigation service advances the DDE stage.", "No document-upload integration is currently invoked from Asset Details."],
        ["Completion, all-assets gate, Next, Back, route, persistence, read-only review, and the current absence of asset-document capture."],
    ),
]


discrepancies = [
    "No dedicated DDE Asset requirements baseline exists locally. The supplied QDE requirements are used only for structure; field and rule content comes from the live screen, repository, and metadata.",
    "The live Pricing Details section displays Other Charges and required RTO / Road Tax, while the reviewed repository displays Motor Insurance Premium and does not map the two live fields in the reviewed save contract.",
    "The live Required Tenure list contains six-month increments from 6 to 84, while the reviewed repository generates every month from 1 through 84.",
    "The live RC Limit Enabled Dealer checkbox is read-only; the reviewed repository renders it editable and maps it to the same stored value used for RC Retention Applicable.",
    "Fetch Price, Vehicle Dedupe, registration report, deal master, vehicle-detail completion, and Hourly viability currently return simulated or fixed responses rather than production service results.",
    "The Used Vehicle Sub-Category values displayed by the UI do not match the restricted values in field metadata and can fail during persistence.",
    "Used-asset save does not mark the line item Asset Details stage complete, while list completion is calculated using a smaller subset than the mandatory Used form.",
    "Any related viability record can satisfy the list-level viability test, including Hire vs Buy or an incomplete assessment.",
    "Hire vs Buy gates New completion for applicable products but is not required for Used completion even when the action is displayed.",
    "Calculation Basis is not stored explicitly; completed viability is inferred from populated fields and can be ambiguous.",
    "Used dedupe success is not reliably invalidated when registration inputs change.",
    "Registration Number and Manufacturing Year & Month do not enforce their displayed formats.",
    "Hire vs Buy permits blank and most negative values; only Hours per Day and Working Days per Month have explicit limits.",
    "Reducing Number of Vehicles does not remove existing sibling assets.",
    "Clone copies identification, type, and purpose but not pricing, dealer, loan, viability, or Hire vs Buy values.",
    "Mixed New, Used, and Refinance assets in one application have no confirmed application-level routing rule.",
    "Product type can be resolved from the last line item and newly created line items do not consistently carry product classification.",
    "Required IRR is functionally a rate but is stored using a currency field type.",
    "No Asset Details-specific document upload control or document-completion gate is present despite documents being requested in the requirements prompt.",
    "The live MMV data exposed Makes with no dependent Models or Models with no dependent Variants; downstream pricing could not be reached through those natural selections.",
    "The New-asset Next action can invoke stage advancement twice through both an event and a direct request, creating a risk of skipping or duplicating navigation.",
    "The Used/Refinance detail component does not correctly consume the completed-stage read-only flag, so a past-stage asset may remain editable.",
    "Used-vehicle integration responses can place display labels into master-code fields without consistently retaining the corresponding display names.",
    "The Used/Refinance component still references removed Make, Model, and Variant accessors in some validation and disabled-state logic.",
    "Manual pricing recalculation excludes insurance, RTO, hypothecation, FASTag, and other charges, and can make On Road Price equal to Ex Showroom Price.",
    "Fetch Price creates or updates asset records before pricing succeeds, so a failed request can leave partial asset records.",
    "Downstream Vehicle Inspection document retrieval is not scoped to a specific asset, creating a risk of selecting the wrong document in a multi-asset application.",
    "Automated controller tests call the additional-asset operation with an outdated parameter list and may not compile against the current implementation.",
    "Dynamic action-handler failures are logged without a user-facing error, so an action can appear to do nothing.",
]


questions = [
    "What is the approved production source and contract for Fetch Price, Vehicle Dedupe, registration report, deal master, vehicle details, and Hourly viability?",
    "Which Pricing Details fields are authoritative: Other Charges and RTO / Road Tax from the live screen, or Motor Insurance Premium from the repository?",
    "Should Required Tenure be limited to six-month increments or allow every month through the policy maximum?",
    "What are the authoritative Used Vehicle Sub-Category values?",
    "Are RC Retention Applicable and RC Limit Enabled Dealer separate business concepts and fields?",
    "Should Hire vs Buy be mandatory for Used and Refinance assets for applicable products?",
    "Should Used save mark Asset Details complete and should list completion enforce every mandatory Used field?",
    "Must Calculation Basis be stored explicitly?",
    "What registration-number and Manufacturing Year & Month formats must be enforced?",
    "Should all money, count, distance, and rate inputs reject negative values, and which can accept zero?",
    "What should happen when Number of Vehicles is reduced after siblings have been created?",
    "Should Clone copy pricing, dealer, loan, viability, and Hire vs Buy details?",
    "Can one application contain mixed New, Used, and Refinance assets; if yes, how is the next DDE step selected?",
    "Which Asset Details documents are required, where should they be uploaded, and which documents block completion?",
    "What timeout, retry-count, and manual fallback policy applies to each external integration?",
    "Are the current dealer, State, and District lists temporary or approved production reference data?",
    "Should completed Asset Details remain read-only, or may Sales Officers reopen and revise them before approval?",
    "What should the user do when a valid Make has no Model or a valid Model has no Variant in the master?",
]


def bullets(items: list[str]) -> str:
    return "\n".join(f"- {item}" for item in items)


def acceptance_text(item: dict) -> str:
    lines = ["Scenarios"]
    for label, criteria in item["scenarios"]:
        lines.append(label)
        lines.extend(f"  • {criterion}" for criterion in criteria)
    lines.extend(["", "Assumptions"])
    lines.extend(f"• {entry}" for entry in item["assumptions"])
    lines.extend(["", "Integrations"])
    lines.extend(f"• {entry}" for entry in item["integrations"])
    lines.extend(["", "In Scope"])
    lines.extend(f"• {entry}" for entry in item["scope"])
    return "\n".join(lines)


def story_markdown(item: dict) -> str:
    lines = [
        f"### {item['id']}: {item['title']}",
        "",
        f"**User Story Description:** {item['description']}",
        "",
        "**Acceptance Criteria:**",
        "",
    ]
    for label, criteria in item["scenarios"]:
        lines.append(f"- **{label}:**")
        lines.extend(f"  - {criterion}" for criterion in criteria)
    lines.extend(
        [
            "",
            "**Assumptions:**",
            "",
            bullets(item["assumptions"]),
            "",
            "**Integrations:**",
            "",
            bullets(item["integrations"]),
            "",
            "**In Scope:**",
            "",
            bullets(item["scope"]),
            "",
        ]
    )
    return "\n".join(lines)


def build_markdown() -> str:
    headers = [
        "Field Name",
        "Object",
        "Field Type",
        "Drop down values",
        "Display Stage",
        "Display Sub Stage",
        "Display Section",
        "Validation",
        "Error Message",
        "Controlling field",
        "Controlling Expression",
    ]
    table = [
        "| " + " | ".join(headers) + " |",
        "|" + "|".join("---" for _ in headers) + "|",
    ]
    for row in fields:
        safe = [str(value or "").replace("|", "\\|").replace("\n", " ") for value in row]
        table.append("| " + " | ".join(safe) + " |")

    story_text = "\n".join(story_markdown(item) for item in stories)
    return f"""# CV DDE Asset Details Requirements

## Document Control

- **Stage:** DDE
- **Sub Stage:** Asset Details
- **Page/Screen Name:** CV DDE Asset Details
- **Page Key:** DDE-ASSET-CV
- **Persona:** Sales Officer
- **Applicability used for Excel:** CE = Yes, CV = Yes, Auto = Yes
- **Workbook update mode:** REPLACE_TARGET_PAGE_ROWS

## Scope and Evidence

No dedicated DDE Asset requirements Markdown was available or required. The existing Commercial Vehicle QDE requirements are used only as the document-structure reference. Functional content is derived from the open Salesforce New Commercial Vehicle application, the reviewed repository, Salesforce metadata, and the blank target workbook sheets.

The live application was used to exercise Add Asset, product-filtered Make/Model/Variant behavior, successful Fetch Price, live Pricing Details, State-to-District dependency, dealer selection, loan fields, six-month tenure options, Hourly and KMS viability branches, and Hire vs Buy. The current application is a New vehicle application; Used and Refinance behavior is documented from the dedicated implementation and metadata and is identified for confirmation where live execution was not possible.

## Observed Screen Structure

- Asset Details opens as a multi-asset list with Add, Edit/View, Clone for New assets, Delete, status, type counts, and pagination.
- New assets progress through Vehicle Selection, Fetch Price, Pricing Details, Dealer & Registration, Loan Information, Viability, and applicable Hire vs Buy.
- Used and Refinance assets progress through Vehicle Dedupe, sequential vehicle-information retrieval, Vehicle Details, Dealer & Registration, Loan Information, Viability, and then Valuation.
- New Vehicle Selection requires master-backed Make, Model, and Variant. The live master contained dependency gaps that prevented natural selection for several valid-looking options.
- No Asset Details-specific document-upload section is present on the reviewed screen or in the reviewed Asset Details components.
- Completed assets are reviewed from the list and later-stage visits are intended to be read-only.

## Excel Field Details Table

{chr(10).join(table)}

## Validation and Dependency Matrix

| Rule | Functional behavior |
|---|---|
| Asset profile | New opens the pricing journey; Used and Refinance open registration dedupe and retrieval. |
| MMV dependency | Product Type → Make → Model → Variant; changing an upstream selection clears dependent values and results. |
| Fetch Price gate | Make, Model, Variant, and a whole Number of Vehicles of at least 1 are required. |
| Pricing calculation | Ex Showroom Price = Basic Price − Discount + GST. The approved treatment of live Other Charges and RTO / Road Tax requires alignment. |
| Dealer dependency | Required pricing reveals Dealer & Registration; State controls District. |
| Loan dependency | Required dealer, State, and District reveal Loan Information. |
| New completion | Required fields, Fetch Price, saved Viability, and applicable saved Hire vs Buy are required. |
| Used completion | Dedupe, retrieval, mandatory vehicle details, dealer, loan, and saved Viability are intended to be required. |
| Hourly viability | Hourly operational income and expense values drive surplus and viability ratio. |
| KMS viability | Route, trips, kilometres, rate, and expenses drive surplus and viability ratio. |
| Hire vs Buy | Hire charges are compared with fuel, staffing, taxes, maintenance, and EMI ownership costs. |
| Multi-asset | Vehicle count can create siblings; New save synchronizes mapped details to existing siblings. |
| Navigation | All assets must be complete; New routes to Loan Terms and Used/Refinance routes to Valuation. |

## User Stories

{story_text}
## Integration Summary

### Vehicle Master Lookup

- **Trigger:** Make, Model, and Variant initialization or an upstream selection change.
- **Prerequisites:** Product classification; Make for Model; Make and Model for Variant.
- **Success:** Active compatible master values are displayed and can be selected.
- **Failure / Timeout:** No explicit timeout or error contract is implemented; the list can appear empty.
- **Retry:** Reopen or search the picker again.
- **Gate:** Missing master selections block Fetch Price.
- **Manual entry:** Not supported.

### Fetch Price

- **Trigger:** Fetch Price after complete vehicle selection and count.
- **Prerequisites:** Make, Model, Variant, valid Number of Vehicles.
- **Success:** Pricing is displayed and downstream sections can be completed.
- **Failure / Timeout:** An error is shown and pricing remains incomplete.
- **Retry:** Fetch Price becomes available again.
- **Gate:** Failure blocks pricing, dealer, loan, viability, and completion.
- **Manual entry:** Not available before a successful response.
- **Current status:** Simulated fixed response.

### Used Vehicle Dedupe

- **Trigger:** Vehicle Dedupe after complete mandatory registration inputs.
- **Prerequisites:** Registration Format, Vehicle Sub-Category, Registration Number.
- **Success:** A positive proceed result enables Get Vehicle Details.
- **Failure / Timeout:** Progress shows the failed step and later actions remain blocked.
- **Retry:** Reruns dedupe.
- **Gate:** A non-proceed or failure blocks retrieval and Next.
- **Manual entry:** No bypass.
- **Current status:** Simulated positive response.

### Registration Report, Vehicle Master, and Vehicle Details

- **Trigger:** Get Vehicle Details after successful dedupe.
- **Prerequisites:** Successful dedupe and unchanged registration inputs.
- **Success:** The three steps complete in sequence and populate vehicle and ownership details.
- **Failure / Timeout:** The sequence stops at the failed step.
- **Retry:** Restarts the full sequence.
- **Gate:** Failure keeps Vehicle Details hidden.
- **Manual entry:** No pre-success fallback.
- **Current status:** Simulated fixed responses.

### Viability

- **Trigger:** Viability Check after loan inputs are complete.
- **Prerequisites:** Active Calculation Basis and all mandatory fields for that branch.
- **Success:** Calculated values and the viability result are saved for the asset.
- **Failure / Timeout:** Hourly processing failure leaves viability incomplete.
- **Retry:** Reruns Hourly processing and save.
- **Gate:** Unsaved viability blocks completion.
- **Manual entry:** KMS calculates internally; Hourly has no service-failure bypass.
- **Current status:** Hourly processing is simulated.

### Hire vs Buy

- **Trigger:** Save within Hire vs Buy.
- **Prerequisites:** Applicable product and complete loan section.
- **Success:** Comparison values are calculated and stored.
- **Failure:** Save error leaves completion unchanged.
- **Retry:** Correct values and Save again.
- **Gate:** Required for New SCV, LCV, MHCV, and CE in the current flow.
- **Manual entry:** Editable comparison inputs are entered directly.

## Implementation Discrepancies Requiring Resolution

{bullets(discrepancies)}

## Clarifying Questions

{bullets(questions)}
"""


def copy_cell(source, target) -> None:
    target.value = source.value
    if source.has_style:
        target._style = copy.copy(source._style)
    target.font = copy.copy(source.font)
    target.fill = copy.copy(source.fill)
    target.border = copy.copy(source.border)
    target.alignment = copy.copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy.copy(source.protection)


def clear_sheet(ws, max_col: int) -> None:
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))
    for row in ws.iter_rows(
        min_row=1,
        max_row=max(ws.max_row, 1),
        min_col=1,
        max_col=max(max_col, ws.max_column),
    ):
        for cell in row:
            cell.value = None
            cell._style = StyleArray()


def copy_template_layout(source, target, header_rows: int, max_col: int) -> None:
    clear_sheet(target, max_col)
    for row_number in range(1, header_rows + 1):
        for column in range(1, max_col + 1):
            copy_cell(source.cell(row_number, column), target.cell(row_number, column))
        target.row_dimensions[row_number].height = source.row_dimensions[row_number].height
    for merged_range in source.merged_cells.ranges:
        if merged_range.max_row <= header_rows and merged_range.max_col <= max_col:
            target.merge_cells(str(merged_range))
    for column in range(1, max_col + 1):
        letter = openpyxl.utils.get_column_letter(column)
        target.column_dimensions[letter].width = source.column_dimensions[letter].width
        target.column_dimensions[letter].hidden = source.column_dimensions[letter].hidden
    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines


def populate_workbook() -> None:
    workbook = openpyxl.load_workbook(XLSX_PATH)
    field_template_name = "DDE_BD_Fields" if "DDE_BD_Fields" in workbook.sheetnames else "QDE_BD_Fields"
    story_template_name = "DDE_BD_US" if "DDE_BD_US" in workbook.sheetnames else "QDE_BD_US"
    field_template = workbook[field_template_name]
    story_template = workbook[story_template_name]
    field_ws = workbook["DDE_ASSET_Fields"]
    story_ws = workbook["DDE_ASSET_US"]

    populated_field_rows = any(
        field_ws.cell(row, 1).value not in (None, "")
        for row in range(2, field_ws.max_row + 1)
    )
    populated_story_rows = any(
        story_ws.cell(row, 1).value not in (None, "")
        for row in range(3, story_ws.max_row + 1)
    )
    if populated_field_rows or populated_story_rows:
        raise RuntimeError("Target DDE Asset worksheets contain populated rows; refusing to overwrite.")

    copy_template_layout(field_template, field_ws, 1, 11)
    for row_number, values in enumerate(fields, start=2):
        for column, value in enumerate(values, start=1):
            target = field_ws.cell(row_number, column, value)
            source = field_template.cell(2, column)
            if source.has_style:
                target._style = copy.copy(source._style)
            target.alignment = copy.copy(source.alignment)
        field_ws.row_dimensions[row_number].height = 84
    field_ws.auto_filter.ref = f"A1:K{len(fields) + 1}"

    copy_template_layout(story_template, story_ws, 2, 14)
    for row_number, item in enumerate(stories, start=3):
        assumptions = "\n".join(f"• {entry}" for entry in item["assumptions"])
        integrations = "\n".join(f"• {entry}" for entry in item["integrations"])
        values = [
            item["id"],
            None,
            item["title"],
            item["description"],
            None,
            acceptance_text(item),
            STAGE,
            PERSONA,
            assumptions,
            integrations,
            None,
            "Yes",
            "Yes",
            "Yes",
        ]
        for column, value in enumerate(values, start=1):
            target = story_ws.cell(row_number, column, value)
            source = story_template.cell(3, column)
            if source.has_style:
                target._style = copy.copy(source._style)
            target.alignment = copy.copy(source.alignment)
        story_ws.row_dimensions[row_number].height = 360
    story_ws.auto_filter.ref = f"A2:N{len(stories) + 2}"

    temporary = XLSX_PATH.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, XLSX_PATH)


def verify() -> None:
    workbook = openpyxl.load_workbook(XLSX_PATH, data_only=False)
    field_ws = workbook["DDE_ASSET_Fields"]
    story_ws = workbook["DDE_ASSET_US"]
    assert field_ws.max_row == len(fields) + 1
    assert story_ws.max_row == len(stories) + 2
    assert [field_ws.cell(row, 1).value for row in range(2, len(fields) + 2)] == [
        row[0] for row in fields
    ]
    assert [story_ws.cell(row, 1).value for row in range(3, len(stories) + 3)] == [
        item["id"] for item in stories
    ]
    for column in (12, 13, 14):
        assert all(
            story_ws.cell(row, column).value == "Yes"
            for row in range(3, len(stories) + 3)
        )

    markdown = MD_PATH.read_text(encoding="utf-8")
    assert markdown.count("### US-DDE-ASSET-CV-") == len(stories)
    assert "Out of Scope" not in markdown
    story_section = markdown.split("## User Stories", 1)[1].split("## Integration Summary", 1)[0]
    assert "__c" not in story_section
    assert "Apex" not in story_section
    assert "LWC" not in story_section


if __name__ == "__main__":
    MD_PATH.write_text(build_markdown(), encoding="utf-8")
    populate_workbook()
    verify()
    print(f"Wrote {MD_PATH.name} with {len(fields)} field rows and {len(stories)} user stories.")
    print("Populated DDE_ASSET_Fields and DDE_ASSET_US in User Stories_CE.xlsx.")
