from __future__ import annotations

import hashlib
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "User stories_Basic_QDE.xlsx"
OUTPUT = ROOT / "User stories_Basic_QDE_Final.xlsx"

FIELDS_SHEET = "QDE_BD_Fields"
STORIES_SHEET = "QDE_BD_US"
DELTA_SHEET = "Delta"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
CRITICAL_FILL = PatternFill("solid", fgColor="F4CCCC")
SUCCESS_FILL = PatternFill("solid", fgColor="D9EAD3")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9E1F2"),
    right=Side(style="thin", color="D9E1F2"),
    top=Side(style="thin", color="D9E1F2"),
    bottom=Side(style="thin", color="D9E1F2"),
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def last_nonempty_row(sheet) -> int:
    for row_number in range(sheet.max_row, 0, -1):
        if any(sheet.cell(row_number, column).value is not None for column in range(1, sheet.max_column + 1)):
            return row_number
    return 1


def find_row(sheet, column: int, expected: str) -> int:
    for row_number in range(1, last_nonempty_row(sheet) + 1):
        if str(sheet.cell(row_number, column).value or "").strip() == expected:
            return row_number
    raise ValueError(f"Could not find {expected!r} in {sheet.title} column {column}")


def copy_row_format(sheet, source_row: int, target_row: int, max_column: int) -> None:
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for column in range(1, max_column + 1):
        source_cell = sheet.cell(source_row, column)
        target_cell = sheet.cell(target_row, column)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)


def trim_blank_rows(sheet) -> None:
    final_row = last_nonempty_row(sheet)
    if sheet.max_row > final_row:
        sheet.delete_rows(final_row + 1, sheet.max_row - final_row)


def set_row_values(sheet, row_number: int, values: list[object]) -> None:
    for column, value in enumerate(values, 1):
        sheet.cell(row_number, column).value = value


def format_data_sheet(sheet, max_column: int) -> None:
    final_row = last_nonempty_row(sheet)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(max_column)}{final_row}"
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2, max_row=final_row, max_col=max_column):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def create_audit_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list[object]],
    widths: list[int],
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER

    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32


def reconcile_fields(sheet) -> None:
    entity_row = find_row(sheet, 1, "Applicant Type (Entity Type)")
    entity_values = [
        value.strip()
        for value in str(sheet.cell(entity_row, 4).value or "").splitlines()
        if value.strip() and value.strip().casefold() != "proprietorship"
    ]
    if len(entity_values) == 9:
        sheet.cell(entity_row, 4).value = "\n".join(entity_values)
    else:
        raise ValueError("Entity Type delta did not produce the expected nine baseline values")

    gender_row = find_row(sheet, 1, "Gender")
    sheet.delete_rows(gender_row, 1)

    incorporation_row = find_row(sheet, 1, "Date of Incorporation")
    age_row = incorporation_row + 1
    sheet.insert_rows(age_row, 1)
    copy_row_format(sheet, age_row + 1, age_row, 11)
    set_row_values(
        sheet,
        age_row,
        [
            "Age",
            "Derived from Applicant__c.Date_of_Birth__c; no Age field exists in repository metadata",
            "Auto-calculated display",
            None,
            "QDE",
            "Basic Details",
            "Document Verification",
            "Optional. Automatically calculated from Date of Birth for an Individual or Date of Incorporation for a Non-Individual. The approved delta does not define the calculation formula, rounding, or persistence.",
            None,
            "Customer Type + Date of Birth / Date of Incorporation",
            "For Individual, calculate from Date of Birth. For Non-Individual, calculate from Date of Incorporation. Leave blank when the applicable source date is blank.",
        ],
    )

    trim_blank_rows(sheet)
    sheet.title = "Final Fields"
    format_data_sheet(sheet, 11)
    widths = [28, 45, 24, 48, 13, 18, 28, 70, 45, 38, 58]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def reconcile_stories(sheet) -> None:
    application_story_row = find_row(sheet, 1, "US-QDE-BASIC-CV-002")
    application_criteria = str(sheet.cell(application_story_row, 6).value or "")
    entity_delta_scenario = (
        "\nDelta Scenario — Entity Type options\n"
        "  • Given Customer Type is Non-Individual\n"
        "  • When Applicant Type (Entity Type) options are displayed\n"
        "  • Then Proprietorship is not available for selection\n"
    )
    if "Then Proprietorship is not available for selection" not in application_criteria:
        sheet.cell(application_story_row, 6).value = application_criteria.rstrip() + "\n" + entity_delta_scenario

    form60_story_row = find_row(sheet, 1, "US-QDE-BASIC-CV-007")
    form60_criteria = str(sheet.cell(form60_story_row, 6).value or "")
    replacements = {
        "Then Gender, First Name, Last Name, and Form 60 upload are displayed": (
            "Then First Name, Last Name, and Form 60 upload are displayed"
        ),
        "Given Form 60 upload, Gender, First Name, Last Name, Date of Birth, or PAN Acknowledgement Number is missing": (
            "Given Form 60 upload, First Name, Last Name, Date of Birth, or PAN Acknowledgement Number is missing"
        ),
    }
    for old_text, new_text in replacements.items():
        if old_text not in form60_criteria:
            raise ValueError(f"Expected Gender reference is missing from US-QDE-BASIC-CV-007: {old_text}")
        form60_criteria = form60_criteria.replace(old_text, new_text)
    sheet.cell(form60_story_row, 6).value = form60_criteria

    final_story_row = max(
        row_number
        for row_number in range(3, last_nonempty_row(sheet) + 1)
        if str(sheet.cell(row_number, 1).value or "").startswith("US-")
    )
    age_story_row = final_story_row + 1
    copy_row_format(sheet, final_story_row, age_story_row, 14)
    sheet.row_dimensions[age_story_row].height = 210
    age_acceptance_criteria = (
        "Scenarios\n"
        "Positive Scenario — Individual\n"
        "  • Given a valid Date of Birth is available for an Individual\n"
        "  • When Document Verification displays the identity details\n"
        "  • Then Age is automatically calculated from Date of Birth and displayed\n"
        "Positive Scenario — Non-Individual\n"
        "  • Given a valid Date of Incorporation is available for a Non-Individual\n"
        "  • When Document Verification displays the entity details\n"
        "  • Then Age is automatically calculated from Date of Incorporation and displayed\n"
        "Optional-field Scenario\n"
        "  • Given the applicable source date is blank\n"
        "  • When Age is evaluated\n"
        "  • Then Age remains blank and does not independently block section completion\n"
        "Recalculation Scenario\n"
        "  • Given the applicable source date changes before identity verification is locked\n"
        "  • When the changed date is accepted\n"
        "  • Then Age is recalculated from the changed source date\n"
        "\nAssumptions\n"
        "• The approved delta defines Age as an automatically calculated, non-mandatory screen element.\n"
        "• Calculation formula, rounding, persistence, and reporting treatment remain unresolved.\n"
        "\nDependencies\n"
        "• Date of Birth or Date of Incorporation must be available before Age can be calculated."
    )
    set_row_values(
        sheet,
        age_story_row,
        [
            "US-QDE-BASIC-CV-010",
            None,
            "Display Automatically Calculated Age",
            "As a Sales Officer, I want Age to be calculated from the applicable birth or incorporation date so that I can view the applicant or entity age without entering it manually.",
            None,
            age_acceptance_criteria,
            "QDE",
            "Sales Officer",
            "• Age is an automatically calculated, non-mandatory screen element.",
            "• Date of Birth or Date of Incorporation must be available.\n• Persistence and reporting treatment require a business and architecture decision.",
            None,
            "Yes",
            "Yes",
            "Yes",
        ],
    )

    trim_blank_rows(sheet)
    sheet.title = "Final User Stories"
    format_data_sheet(sheet, 14)
    widths = [24, 12, 38, 58, 18, 95, 12, 22, 50, 55, 34, 12, 12, 12]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {OUTPUT}")

    source_hash_before = sha256(SOURCE)
    workbook = load_workbook(SOURCE)
    expected_sheets = {FIELDS_SHEET, STORIES_SHEET, DELTA_SHEET}
    if not expected_sheets.issubset(workbook.sheetnames):
        raise ValueError(f"Expected sheets {sorted(expected_sheets)}, found {workbook.sheetnames}")

    reconcile_fields(workbook[FIELDS_SHEET])
    reconcile_stories(workbook[STORIES_SHEET])
    workbook.remove(workbook[DELTA_SHEET])

    delta_log_rows = [
        [
            "DELTA-001",
            18,
            "Remove the field",
            "Gender",
            "Remove Gender from Basic Details.",
            "Removed the Gender row from Final Fields.",
            "Removed Gender from the positive and negative Form 60 scenarios in US-QDE-BASIC-CV-007.",
            "Applicant__c.Gender__c and Document_CheckList__c.Gender__c still exist; the current LWC still renders and validates Gender.",
            "Applied to final requirements; implementation discrepancy logged.",
        ],
        [
            "DELTA-002",
            19,
            "Add a new field",
            "Age",
            "Add non-mandatory Age, automatically calculated from Date of Birth or Date of Incorporation, applicable to CV/CE and Auto.",
            "Added Age after Date of Incorporation as a derived, non-mandatory Document Verification field.",
            "Added US-QDE-BASIC-CV-010 with Individual, Non-Individual, optional-field, and recalculation scenarios.",
            "No Applicant Age field exists in repository metadata. Applicant__c.Date_of_Birth__c exists and is also the baseline mapping for Date of Incorporation.",
            "Applied as derived display; persistence discrepancy logged.",
        ],
        [
            "DELTA-003",
            20,
            "Update dropdown",
            "Applicant__c.Entity_Type__c",
            "Remove Proprietorship from the Entity Type dropdown.",
            "Removed Proprietorship from the Final Fields dropdown values while preserving the other baseline values.",
            "Added an Entity Type delta scenario to US-QDE-BASIC-CV-002.",
            "Entity_Type__c exists, but repository metadata still defines Proprietorship as an active value.",
            "Applied to final requirements; metadata discrepancy logged.",
        ],
    ]
    create_audit_sheet(
        workbook,
        "Applied Delta Log",
        [
            "Delta ID",
            "Source Row",
            "Change Type",
            "Item",
            "Approved Delta",
            "Final Fields Action",
            "Final User Stories Action",
            "Metadata / Code Check",
            "Status",
        ],
        delta_log_rows,
        [14, 12, 20, 30, 55, 55, 65, 75, 38],
    )
    for row in workbook["Applied Delta Log"].iter_rows(min_row=2):
        for cell in row:
            cell.fill = SUCCESS_FILL

    discrepancy_rows = [
        [
            "DISC-001",
            "High",
            "Age persistence and calculation",
            "Delta row 19",
            "The delta identifies SF Object as Applicant and says Age is auto-calculated, but no Age field exists under Applicant__c metadata. Formula, rounding, persistence, and reporting behavior are not defined.",
            "Final Fields treats Age as a derived display and US-QDE-BASIC-CV-010 avoids inventing a formula.",
            "Confirm whether Age is display-only or persisted; if persisted, approve the real field API name, formula, classification, and reporting behavior.",
            "Open",
        ],
        [
            "DISC-002",
            "High",
            "Gender removal implementation gap",
            "Delta row 18 vs repository",
            "The approved delta removes Gender, while Applicant__c.Gender__c and Document_CheckList__c.Gender__c remain in metadata and applicantBasicDetailsCapture still displays and requires Gender on the Individual Form 60 path.",
            "Gender is removed from Final Fields and US-QDE-BASIC-CV-007 only.",
            "Confirm that removal is screen-specific, then update the UI and completion logic without deleting shared KYC fields unless separately approved.",
            "Open",
        ],
        [
            "DISC-003",
            "High",
            "Entity Type values",
            "Delta row 20, baseline, and metadata",
            "The baseline dropdown values do not match the active Entity_Type__c metadata values, and Proprietorship remains active in metadata despite the approved removal.",
            "Final Fields preserves baseline values except Proprietorship, as instructed.",
            "Approve the complete post-delta Entity Type value set and align metadata/UI. Do not deactivate shared values until impact outside Basic Details is assessed.",
            "Open",
        ],
        [
            "DISC-004",
            "Medium",
            "Date of Incorporation mapping",
            "Baseline Final Fields and repository",
            "Date of Incorporation is mapped to Applicant__c.Date_of_Birth__c and carries 18–80-year person-age errors. The new Age requirement depends on that date for Non-Individual applicants.",
            "The existing mapping is preserved and explicitly identified as the source for derived Age.",
            "Confirm whether Non-Individual incorporation date should continue using Date_of_Birth__c and whether person-age limits apply to entity age.",
            "Open",
        ],
    ]
    create_audit_sheet(
        workbook,
        "Unresolved Discrepancies",
        [
            "Discrepancy ID",
            "Severity",
            "Area",
            "Sources",
            "Evidence",
            "Current Treatment",
            "Decision Required",
            "Status",
        ],
        discrepancy_rows,
        [16, 12, 32, 32, 80, 60, 78, 14],
    )
    discrepancy_sheet = workbook["Unresolved Discrepancies"]
    for row in discrepancy_sheet.iter_rows(min_row=2):
        fill = CRITICAL_FILL if row[1].value == "High" else WARNING_FILL
        for cell in row:
            cell.fill = fill

    workbook.properties.title = "Basic QDE Consolidated Fields and User Stories"
    workbook.properties.subject = "Baseline requirements reconciled with approved Delta sheet"
    workbook.properties.description = (
        "Generated from User stories_Basic_QDE.xlsx. Source workbook preserved. "
        "See Applied Delta Log and Unresolved Discrepancies."
    )
    workbook.save(OUTPUT)

    if sha256(SOURCE) != source_hash_before:
        raise RuntimeError("Source workbook changed unexpectedly")

    verification = load_workbook(OUTPUT, read_only=True, data_only=False)
    if verification.sheetnames != [
        "Final Fields",
        "Final User Stories",
        "Applied Delta Log",
        "Unresolved Discrepancies",
    ]:
        raise AssertionError(verification.sheetnames)
    fields = verification["Final Fields"]
    stories = verification["Final User Stories"]
    field_names = [fields.cell(row, 1).value for row in range(2, fields.max_row + 1)]
    story_ids = [stories.cell(row, 1).value for row in range(3, stories.max_row + 1)]
    if "Gender" in field_names or "Age" not in field_names:
        raise AssertionError(field_names)
    if "US-QDE-BASIC-CV-010" not in story_ids:
        raise AssertionError(story_ids)
    all_story_text = "\n".join(
        str(stories.cell(row, column).value or "")
        for row in range(3, stories.max_row + 1)
        for column in range(1, 15)
    )
    if "Gender" in all_story_text:
        raise AssertionError("Gender remains in Final User Stories")
    if "Then Proprietorship is not available for selection" not in all_story_text:
        raise AssertionError("Entity Type delta scenario is missing")
    verification.close()

    print(OUTPUT)


if __name__ == "__main__":
    main()
