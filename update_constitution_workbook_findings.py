from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill


path = Path(__file__).resolve().parent / "Moga_Dev_Constitution_Mandatory_Documents.xlsx"
wb = load_workbook(path)

readme = wb["Read Me"]
readme.append(
    [
        "Government / Form 60 defect",
        "Government metadata requires Form 60, but KYCDetailsController excludes Form 60 "
        "from entity-additional rows and does not otherwise pass it to the mandatory evaluator. "
        "The rule therefore appears impossible to satisfy.",
    ]
)
readme.append(
    [
        "Unknown product behavior",
        "Products that do not normalize to CV or PV normally find no constitution metadata row, "
        "producing an empty constitution-required document list.",
    ]
)
readme.append(
    [
        "Form 60 code evidence",
        "force-app/main/default/classes/KYCDetailsController.cls:29-34, 620-650, 820-899, 1260-1277",
    ]
)
for row in readme.iter_rows(min_row=readme.max_row - 2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

quality = wb["Data Quality"]
quality.append(
    [
        "Critical",
        "Government Form 60 cannot satisfy KYC evaluator",
        "CV/PV",
        80,
        "Government",
        "Required_Documents__c includes Form 60",
        "Form 60 is loaded and maintained by Basic Details, but the KYC entity bucketing "
        "explicitly excludes it. The evaluator receives PAN plus CIN/GST/address/additional rows, "
        "not Form 60, and only PAN is force-treated as collected. Government KYC therefore "
        "appears permanently incomplete.",
    ]
)
quality.append(
    [
        "Warning",
        "Unknown product yields no constitution rule",
        "Other",
        "",
        "",
        "Application__c.Product__c",
        "The product normalizer returns unrecognized values unchanged. Unless matching metadata "
        "exists, the lookup returns no row and no constitution documents are enforced.",
    ]
)

critical_fill = PatternFill("solid", fgColor="F4CCCC")
warning_fill = PatternFill("solid", fgColor="FFF2CC")
for row_number in (quality.max_row - 1, quality.max_row):
    fill = critical_fill if quality.cell(row_number, 1).value == "Critical" else warning_fill
    for cell in quality[row_number]:
        cell.fill = fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)

quality.auto_filter.ref = quality.dimensions
wb.save(path)

check = load_workbook(path, read_only=True, data_only=True)
assert check["Data Quality"].max_row - 1 == 13
assert check["Read Me"].max_row - 1 == 24
check.close()

print(path)
